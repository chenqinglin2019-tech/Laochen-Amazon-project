"""Business view of the quick screen and table I/O for the 禁限售 Skill. Python 3.9+, standard library only
(openpyxl is required for .xlsx, to preserve ordinary workbook content and formatting).

Marks, most severe first: 禁售 › 需平台审核 › 物流运输强限制 › 需要付费认证 › 需工厂提供证书 › 需单独描述 › 未发现限制. A rough
screen always decides: what the text leaves open is filled with the category's common configuration (typical blade
length, laser power, e-mobility power/speed; the agent's "assume" note for traits) and the reason column says so. Paid = certification, third-party test report, registration, license, paid local representative; insurance is
not counted as paid. Account-level obligations (general_obligations) never enter a row; they are listed once."""
import csv
import html.parser
import io
import re
import time
import urllib.request
import zipfile
import pathlib
import tempfile
import csc_runtime as runtime
from xml.etree import ElementTree as ET

MARKS = ("禁售", "需平台审核", "物流运输强限制", "需要付费认证", "需工厂提供证书", "需单独描述", "未发现限制")
PRIMARY_FORMS = ("liquid", "paste_gel", "other", "unknown")
LOGISTICS_NOTE = "按液体/膏体运输选品口径标记；不等于危险品或法定禁运，具体渠道要求按实际条件核验"


def category_error(cat):
    if cat is None:
        return None
    if not isinstance(cat, dict):
        return "分类必须是 JSON 对象"
    form = cat.get("primary_form")
    if form is not None and form not in PRIMARY_FORMS:
        return "primary_form 只能是 liquid / paste_gel / other / unknown"
    evidence = cat.get("form_evidence")
    if evidence is not None and not isinstance(evidence, str):
        return "form_evidence 必须是字符串"
    if form in PRIMARY_FORMS[:3] and not (evidence or "").strip():
        return "已知主体形态必须填写 form_evidence（原文片段或有依据的语义解释）"
    return None


def category_complete(cat):
    return (isinstance(cat, dict) and cat.get("primary_form") in PRIMARY_FORMS
            and not category_error(cat))


def finalize_mark(result, cat):
    """Apply the cross-market selection policy, preserving the underlying compliance result and costs."""
    c = dict(result)
    c["compliance_mark"] = c["mark"]
    c["classification_complete"] = category_complete(cat)
    c["primary_form"] = (cat or {}).get("primary_form") or "unknown"
    c["form_evidence"] = (cat or {}).get("form_evidence") or ""
    c["logistics_restricted"] = c["primary_form"] in ("liquid", "paste_gel") and c["classification_complete"]
    if c["logistics_restricted"]:
        remaining = [item for key in ("paid_items", "factory_items") for item in c.get(key) or []
                     if item not in c["detail"]]
        if remaining:
            c["detail"] += "。其他合规事项：" + "；".join(dict.fromkeys(remaining))
        form = "液体" if c["primary_form"] == "liquid" else "膏霜/凝胶"
        reason = f"物流运输强限制：商品主体为{form}；依据：{c['form_evidence']}；{LOGISTICS_NOTE}"
        if RANK.get(c["mark"], 0) < RANK["物流运输强限制"]:
            c["mark"] = "物流运输强限制"
            c["detail"] = reason + "。原合规要求：" + c["detail"]
        else:
            c["detail"] += "。" + reason
    elif not c["classification_complete"] or c["primary_form"] == "unknown":
        pending = ("商品主体形态尚未完成语义分类，本结果为中间稿；请依据完整资料补充 primary_form 和 form_evidence"
                   if not c["classification_complete"] else "现有资料无法判断商品主体形态，未按非液体处理")
        if c["mark"] not in ("禁售", "需平台审核"):
            c["detail"] = pending + "。原合规初判（" + c["mark"] + "）：" + c["detail"]
            c["mark"] = "无法判断"
        else:
            c["detail"] += "。" + pending
    assume = str((cat or {}).get("assume") or "").strip()
    if assume:
        c["detail"] += "。按常见情形判定：" + assume
    return c
PAID_KINDS = {"certification", "test_report", "registration", "license", "local_representative"}
INSURANCE = re.compile(r"保险|COI|insurance", re.I)
MAX_ITEMS, ITEM_CHARS, CELL_CHARS = 8, 60, 900

# Marketplace by URL host (longest suffix first). Other Amazon hosts are outside the rule pack.
HOST_SITE = [("amazon.com.au", "AU"), ("amazon.co.jp", "JP"), ("amazon.co.uk", "UK"), ("amazon.de", "DE"),
             ("amazon.fr", "FR"), ("amazon.it", "IT"), ("amazon.es", "ES"), ("amazon.ae", "AE"), ("amazon.sa", "SA"),
             ("amazon.com", "US")]
SITE_ALIASES = {"美国": "US", "日本": "JP", "英国": "UK", "GB": "UK", "德国": "DE", "法国": "FR", "意大利": "IT",
                "西班牙": "ES", "阿联酋": "AE", "沙特": "SA", "澳大利亚": "AU", "澳洲": "AU"}


def site_of_url(url):
    """(site, error). site is None when the URL is not an Amazon link of a covered marketplace."""
    m = re.match(r"\s*(?:https?://)?([^/\s?#]+)", str(url or ""), re.I)
    if not m:
        return None, None
    host = m.group(1).lower().split(":")[0]
    for suffix, site in HOST_SITE:
        if host == suffix or host.endswith("." + suffix):
            return site, None
    if "amazon." in host:
        return None, f"站点不在覆盖范围（{host}）"
    if host in ("amzn.to", "a.co", "amzn.eu", "amzn.asia"):
        return None, "短链接看不出站点，请换成完整商品链接"
    return None, None


def norm_site(value):
    v = str(value or "").strip()
    if not v:
        return None
    for k, s in SITE_ALIASES.items():
        if k in v:
            return s
    v = v.upper().replace("站", "").strip()
    return SITE_ALIASES.get(v, v)


def _short(text, n=ITEM_CHARS):
    t = re.sub(r"\s+", " ", str(text or "")).strip()
    return t if len(t) <= n else t[:n - 1] + "…"


def _join(items, cap=MAX_ITEMS):
    seen, out = set(), []
    for x in items:
        if x and x not in seen:
            seen.add(x)
            out.append(x)
    s = "；".join(out[:cap])
    if len(out) > cap:
        s += f"；等共 {len(out)} 项"
    return s


# Numeric thresholds that decide a ban (laser power, blade length, e-mobility power/speed). Keywords cannot see them,
# so they are read from the text; a ban rule that waits on one of them is never silently dropped.
NUMERIC = {"laser.power_mw": "激光功率", "blade.length_cm": "刃长", "power.rated_power_w": "额定功率",
           "mobility.max_speed_kmh": "最高速度", "blade.mechanism": "开合方式"}
# The product kinds each threshold is about: a rule waits on a number only for these products.
NUMERIC_FAMILIES = {"laser.": {"laser_product"}, "blade.": {"knife_blade", "weapon_other"},
                    "power.rated_power_w": {"bicycle", "personal_transporter"},
                    "mobility.": {"bicycle", "personal_transporter"}}
THRESHOLDS = {"laser.power_mw", "blade.length_cm", "power.rated_power_w", "mobility.max_speed_kmh"}
MECHANISM = [(r"switch ?blade|automatic knife|\botf\b|out[- ]the[- ]front", "automatic_switchblade"),
             (r"butterfly knife|balisong", "butterfly"), (r"gravity knife", "gravity"),
             (r"spring[- ]assist", "spring_assisted"),
             (r"fixed[- ]blade|feststehend|lame fixe|lama fissa|hoja fija|chef.?s? knife|kitchen knife|paring knife|santoku|bread knife|steak knife|cleaver|nakiri|boning knife|fillet knife|包丁|菜刀|厨刀", "fixed"),
             (r"liner lock|frame lock|lockback|lock blade|axis lock|locking (?:folding )?(?:knife|blade)", "folding_locking"),
             (r"slip ?joint|non[- ]locking", "folding_manual")]
_NUM = r"(\d+(?:[.,]\d+)?)"


# What the market usually sells when the text gives no number: (value, unit, label). Fixed blades are hunting/outdoor
# knives (kitchen knives are food_contact and excluded by the rules themselves); folding knives are pocket knives.
HIGH_POWER = re.compile(r"high[- ]?power|powerful|long[- ]range|burning|\d{3,}\s*(?:m|mile|meter|metre|ft|feet)\b|"
                        r"大功率|强光|远射|高功率|ハイパワー|hochleistung|puissant|potente", re.I)


def typical_numbers(text, families, mechanism, wanted):
    """{attr: (value, quote)} for thresholds the text leaves open, from the category's usual specification."""
    fams, out = set(families or []), {}
    if "laser.power_mw" in wanted:
        high = HIGH_POWER.search(text or "")
        out["laser.power_mw"] = ((100.0, f"未写功率，文字称「{high.group(0)}」，按大功率（>5 mW）") if high
                                 else (4.0, "未写功率，按常见 <5 mW"))
    if "blade.length_cm" in wanted:
        out["blade.length_cm"] = ((10.0, "未写刃长，固定刃刀常见约 10 cm") if mechanism == "fixed"
                                  else (8.0, "未写刃长，折刀/小刀常见约 8 cm"))
    if "power.rated_power_w" in wanted:
        out["power.rated_power_w"] = ((250.0, "未写功率，电助力自行车常见 250 W") if "bicycle" in fams
                                      else (350.0, "未写功率，电动滑板车/平衡车常见 350 W"))
    if "mobility.max_speed_kmh" in wanted:
        out["mobility.max_speed_kmh"] = (25.0, "未写速度，常见限速 25 km/h")
    return out


def extract_numbers(text):
    """{attr: (value, quote)} read from free text. Conservative: blade length must sit next to "blade"/刃."""
    t = " ".join(str(text or "").split())
    low = t.lower()
    out = {}
    m = re.search(_NUM + r"\s*mw\b", low)
    if m:
        out["laser.power_mw"] = (float(m.group(1).replace(",", ".")), m.group(0))
    unit = {"mm": 0.1, "cm": 1.0, "inch": 2.54, "inches": 2.54, "in": 2.54, '"': 2.54, "″": 2.54, "厘米": 1.0, "公分": 1.0, "英寸": 2.54}
    pat = _NUM + r"\s*(mm|cm|inches|inch|in\b|\"|″|厘米|公分|英寸)"
    # "Chef Knife 8 inch" names the blade too (knife sizes are blade lengths); the blade word wins when both appear
    near = r"(?:blade|刃|刀刃|klinge|lame|lama|hoja|knife|knives|messer|couteau|coltello|cuchillo|包丁|刀)"
    for m in re.finditer(near + r"[^.;,，。；]{0,25}?" + pat + "|" + pat + r"[^.;,，。；]{0,12}?" + near, low):
        num, u = (m.group(1), m.group(2)) if m.group(1) else (m.group(3), m.group(4))
        out["blade.length_cm"] = (round(float(num.replace(",", ".")) * unit[u.strip()], 1), m.group(0))
        break
    m = re.search(_NUM + r"\s*(?:w|watts?)\b(?!h)", low)
    if m:
        out["power.rated_power_w"] = (float(m.group(1).replace(",", ".")), m.group(0))
    for pat, value in MECHANISM:
        m = re.search(pat, low)
        if m:
            out["blade.mechanism"] = (value, m.group(0))
            break
    m = re.search(_NUM + r"\s*(km/h|kmh|kph|mph)\b", low)
    if m:
        v = float(m.group(1).replace(",", "."))
        out["mobility.max_speed_kmh"] = (round(v * (1.609 if m.group(2) == "mph" else 1), 1), m.group(0))
    return out


def numeric_flags(screen):
    """Ban/approval rules decided by a number: (decided, waiting). decided = triggered by a number read from the
    text; waiting = would decide once a still-unknown number is known."""
    nums = screen.get("numbers") or {}
    decided, waiting = [], []
    for c in (screen.get("match") or {}).get("candidates") or []:
        if c.get("severity") not in ("prohibited", "approval_required") or not c.get("rule_id"):
            continue
        if c["kind"] == "triggered" and nums and c.get("_numeric_decisive"):
            decided.append(c)
        elif c["kind"] == "possible" and c.get("_numeric_rule"):
            waiting.append(c)
    return decided, waiting


RANK = {"禁售": 6, "需平台审核": 5, "物流运输强限制": 4, "需要付费认证": 3, "需工厂提供证书": 2, "需单独描述": 1, "未发现限制": 0}
FLAG_ZH = {"children": "儿童用品", "battery_lithium": "含锂电池", "button_cell": "含纽扣电池", "wireless": "无线发射",
           "mains": "插电电器"}


def active_flags(cat):
    """Product traits from the agent's classification that the business table prices on their own: the four core
    traits, button/coin cells, and any site-defined extra flags the agent listed under "flags"."""
    cat = cat or {}
    core = [f for f, on in (("children", cat.get("children") is True), ("battery_lithium", cat.get("battery") == "lithium"),
                            ("button_cell", cat.get("battery") == "button"), ("wireless", cat.get("wireless") is True),
                            ("mains", cat.get("mains") is True)) if on]
    return core + [str(f) for f in cat.get("flags") or [] if str(f) not in core]


def _dedupe(items):
    """Same requirement named by two entries ("CPSC 认可实验室检测（ASTM F963…）" / "（按产品适用标准）"): keep the first."""
    seen, out = set(), []
    for i in items:
        key = re.split(r"[（(：:]", i)[0].strip()
        if key not in seen:
            seen.add(key)
            out.append(i)
    return out


def table_mark(table, families, cat, pack):
    """Mark from the site's business table (family x site, plus traits) -> (mark, detail, paid_items, factory_items)."""
    entries = [(pack.families.get(f, {}).get("zh", f).split("（")[0], table["families"][f])
               for f in families if f in table["families"]]
    flags = table.get("flags", {})
    entries += [(flags[f].get("zh") or FLAG_ZH.get(f, f), flags[f]) for f in active_flags(cat) if f in flags]
    if not entries:
        d = table["default"]
        return d["mark"], d.get("note") or "未归入受管品类", [], []
    mark = max((e["mark"] for _, e in entries), key=lambda m: RANK[m])
    at = lambda m: [(n, e) for n, e in entries if e["mark"] == m]  # noqa: E731
    items = lambda rows: [i for _, e in rows for i in e.get("items") or []]  # noqa: E731
    names = lambda rows: "、".join(dict.fromkeys(n for n, _ in rows))  # noqa: E731
    paid, supplier, desc = (_dedupe(items(at(m))) for m in ("需要付费认证", "需工厂提供证书", "需单独描述"))
    tail = lambda d, sup=True, dsc=True: d + (("。另需工厂提供：" + _join(supplier, 4)) if sup and supplier else "") + (  # noqa: E731
        ("。另需描述：" + _join(desc, 3)) if dsc and desc else "")
    if mark in ("禁售", "需平台审核"):
        top = at(mark)
        detail = ("禁售" if mark == "禁售" else "需审核") + f"（{names(top)}）：" + _join([_short(e["why"], 70).rstrip("。") for _, e in top], 2)
        if items(top):
            detail += "。需准备：" + _join(items(top), 5)
        if mark == "需平台审核":
            detail = tail(detail + (("。另需付费认证：" + _join(paid, 5)) if paid else ""), dsc=False)
    elif mark == "需要付费认证":
        detail = tail(f"需要（{names(at(mark))}）：" + _join(paid))
    elif mark == "需工厂提供证书":
        detail = tail(f"向工厂索取（{names(at(mark))}）：" + _join(supplier) + "。工厂没有则须付费检测/认证", sup=False)
    elif mark == "需单独描述":
        detail = f"需描述（{names(at(mark))}）：" + _join(desc)
    else:
        notes = [_short(e.get("note"), 120) for _, e in entries if e.get("note")]
        detail = "无强制认证" + (f"（{names(entries)}）" if entries else "") + ("。" + _join(notes, 2) if notes else "")
    return mark, detail, paid, supplier


def business(screen, table, cat, pack):
    """Business marks for a site with a table: bans/approvals/unverified bans from the rule pack, the rest from the
    table; the heavier of the two wins."""
    tm, td, paid, supplier = table_mark(table, screen.get("families") or [], cat, pack)
    c = classify(screen, pack, fallback=(tm, td, paid, supplier))
    c["paid_items"], c["factory_items"] = paid, supplier
    if RANK[tm] > RANK[c["mark"]]:
        c["mark"], c["detail"] = tm, td
    claim = next((d for d in c["describe"] if d.startswith("删除原文案中的宣称")), None)
    if claim and c["mark"] != "禁售":
        if c["mark"] == "未发现限制":
            c["mark"], c["detail"] = "需单独描述", claim
        elif claim not in c["detail"]:
            c["detail"] += "。注意：" + claim
    if len(c["detail"]) > CELL_CHARS:
        c["detail"] = c["detail"][:CELL_CHARS - 1] + "…"
    return c


def _plain_q(q):
    """"请确认：刃长（cm）（blade.length_cm），可选：a/b" -> "刃长（cm）" for people reading a spreadsheet."""
    q = re.sub(r"^请确认[：:]\s*", "", str(q))
    q = re.sub(r"[，,]\s*可选[：:].*$", "", q)
    return re.sub(r"（[a-z0-9_]+\.[a-z0-9_.]+）", "", q).strip()


def reasons_of(cands):
    return _join([f"{_short(c['topic_zh'], 70)}（{c['rule_id']}）" for c in cands], 3)


def classify(screen, pack=None, fallback=None):
    """One site's quick screen (ce.quick_screen output) -> {"mark", "detail", "paid", "describe", "maybe"}.

    Requirements of rules that only *possibly* apply (an unknown attribute such as "for children") are listed as
    "视情况", not counted. Rules triggered only by a claim in the text (e.g. "antibacterial") ask to drop the claim:
    their requirements are what keeping the claim would cost, so they are not counted as paid either."""
    result, families = screen["result"], screen.get("families")
    applied = result.get("applied") or []
    sev = {a["rule_id"]: a["severity"] for a in applied}
    topic = {a["rule_id"]: a["topic_zh"] for a in applied}
    by = lambda s: [a for a in applied if a["severity"] == s]  # noqa: E731
    kinds = {}
    for c in (screen.get("match") or {}).get("candidates") or []:
        if c.get("rule_id"):
            kinds.setdefault(c["rule_id"], c["kind"])
    strong = {c["rule_id"] for c in (screen.get("match") or {}).get("candidates") or []
              if c.get("rule_id") and (c.get("keyword_hits") or {}).get("strong")}
    possible = {r for r, k in kinds.items() if k == "possible" and r not in strong}
    profile = screen.get("profile") or {}
    text_claims = {c["id"] for c in profile.get("claims") or []}
    claim_driven = set()
    if pack is not None and text_claims:
        from compliance_engine.match import _Eval  # same engine the Skill ships
        ev = _Eval(pack, profile, without_claims=True)
        for r in sev:
            rule = pack.rules.get(r)
            if rule and r not in possible and ev.cond(rule["triggers"]) is False:
                claim_driven.add(r)
    paid, describe, maybe, claim_cost = [], [], [], []
    for q in result.get("requirements") or []:
        rules = [r for r in q.get("rules") or [] if sev.get(r) not in ("prohibited", "info")]
        if not rules:
            continue
        name = _short(q.get("name_zh"))
        if all(r in possible for r in rules):
            maybe.append(name)
            continue
        if all(r in claim_driven or r in possible for r in rules):
            claim_cost.append(name)
            continue
        if q.get("kind") in PAID_KINDS and not INSURANCE.search(q.get("name_zh") or "") and "可选" not in name:
            paid.append(name)
        elif q.get("kind") != "amazon_approval" or not any(sev.get(r) == "approval_required" for r in rules):
            describe.append(name)
    cons = result.get("constraints") or {}
    found = [c.get("claim_zh") for c in cons.get("forbid_claims") or [] if c.get("claim") in text_claims]
    if found:
        line = "删除原文案中的宣称：" + "、".join(dict.fromkeys(found))
        if claim_cost:
            line += "（保留则需：" + _join(claim_cost, 2) + "）"
        describe.insert(0, line)
    describe += [_short(s.get("text_zh"), 80) for s in cons.get("required_statements") or []
                 if s.get("rule") not in possible]
    verdict = result["verdict"]
    reasons = lambda rows: _join([f"{_short(a['topic_zh'], 70)}（{a['rule_id']}）" for a in rows], 5)  # noqa: E731
    sure = {c["rule_id"] for c in (screen.get("match") or {}).get("candidates") or [] if c.get("can_prohibit")}
    # a ban stands when the rule is verified, or when the product plainly meets it (triggered by what the text/agent
    # says it is, or a strong keyword) even though the rule still awaits verification; a ban only hinted at by a
    # weak keyword does not decide the row
    firm = [a for a in by("prohibited") if a["rule_id"] in sure]
    plain = [a for a in by("prohibited") if a["rule_id"] not in sure
             and (kinds.get(a["rule_id"]) == "triggered" or a["rule_id"] in strong)]
    hinted = [a for a in by("prohibited") if a not in firm and a not in plain]
    if verdict == "prohibited" and not firm and not plain:
        verdict = "approval_required" if by("approval_required") else "compliance_required"
    if verdict == "prohibited" and firm:
        mark, detail = "禁售", "禁售原因：" + reasons(firm)
    elif verdict == "prohibited":
        mark, detail = "禁售", "禁售原因（政策依据待核实，按从严处理）：" + reasons(plain)
    elif verdict == "approval_required":
        mark, detail = "需平台审核", "需审核：" + reasons(by("approval_required"))
        extra = fallback[2] if fallback else paid
        if extra:
            detail += "。另需付费认证：" + _join(extra, 5)
        if fallback and fallback[3]:
            detail += "。另需工厂提供：" + _join(fallback[3], 4)
    elif fallback:
        # the site has a business table: below approval level it decides (rule requirement lists do not)
        mark, detail = fallback[0], fallback[1]
    elif paid:
        mark, detail = "需要付费认证", "需要：" + _join(paid)
        if describe:
            detail += "。另需在 Listing/包装中说明：" + _join(describe, 3)
    elif describe:
        rule_topics = [topic[r] for q in result.get("requirements") or [] for r in q.get("rules") or []
                       if r in topic and r not in possible]
        mark, detail = "需单独描述", "需描述：" + _join(describe)
        if rule_topics:
            detail += "。原因：" + _join([_short(t, 50) for t in rule_topics], 3)
    else:
        mark = "未发现限制"
        detail = "本规则库未匹配到限制" + ("" if families else "；未从文字识别出商品类别，建议人工复核")
    decided, waiting = numeric_flags(screen)
    nums = screen.get("numbers") or {}
    typical = screen.get("typical") or {}
    said = "、".join(f"{NUMERIC[a]}「{q}」" for a, (v, q) in nums.items() if a in NUMERIC and a not in typical)
    guessed = "、".join(q for a, (v, q) in typical.items())
    basis = "；".join(x for x in (said and "按文字中的数值：" + said, guessed and "按常见规格：" + guessed) if x)
    rank = {"禁售": 5, "需平台审核": 4}
    for c in decided:
        new = "需平台审核" if c["severity"] != "prohibited" else "禁售"
        if rank.get(new, 0) > rank.get(mark, 0):
            head = "需审核" if new == "需平台审核" else ("禁售原因" if c.get("can_prohibit") else "禁售原因（政策依据待核实，按从严处理）")
            detail = f"{head}（{basis}）：{_short(c['topic_zh'], 90)}（{c['rule_id']}）"
            mark = new
    spared = [c for c in screen.get("spared") or [] if rank.get(mark, 0) < 5]
    if spared and mark != "禁售":
        word = "禁售" if any(c["severity"] == "prohibited" for c in spared) else "需审核"
        detail += f"。{basis or '参数未写明，按常见规格'}，据此判定；若实际超出则{word}：{reasons_of(spared)}"
    if waiting and mark != "禁售":
        # a threshold with no common value (should not happen for the families we know): decide on the rest
        word = "禁售" if any(c["severity"] == "prohibited" for c in waiting) else "需审核"
        detail += f"。参数未写明，按未触发判定；若实际满足则{word}：{reasons_of(waiting)}"
    if hinted and mark != "禁售":
        detail += f"。按常见情形未判禁售；若实际属于以下情形则禁售：{reasons(hinted)}"
    if maybe and mark not in ("禁售", "未发现限制") and not fallback:
        detail += "。视情况（如儿童用品、特定材质）可能另需：" + _join(maybe, 2)
    if len(detail) > CELL_CHARS:
        detail = detail[:CELL_CHARS - 1] + "…"
    return {"mark": mark, "detail": detail, "paid": paid, "describe": describe, "maybe": maybe}


# ---------------------------------------------------------------- product page (best effort)

class _PageParser(html.parser.HTMLParser):
    IDS = {"productTitle": "title", "feature-bullets": "bullets", "productFactsDesktopExpander": "bullets",
           "productDescription": "description"}

    def __init__(self):
        super().__init__()
        self.stack, self.out, self.cur_li = [], {"title": [], "bullets": [], "description": []}, None

    def handle_starttag(self, tag, attrs):
        if tag in ("br", "img", "meta", "link", "input", "hr"):
            return
        field = self.IDS.get(dict(attrs).get("id") or "")
        parent = self.stack[-1] if self.stack else None
        self.stack.append(field or parent)
        if tag == "li" and self.stack[-1] == "bullets":
            self.cur_li = []

    def handle_endtag(self, tag):
        if tag in ("br", "img", "meta", "link", "input", "hr") or not self.stack:
            return
        if tag == "li" and self.cur_li is not None:
            text = " ".join("".join(self.cur_li).split())
            if text:
                self.out["bullets"].append(text)
            self.cur_li = None
        self.stack.pop()

    def handle_data(self, data):
        field = self.stack[-1] if self.stack else None
        if field == "bullets" and self.cur_li is not None:
            self.cur_li.append(data)
        elif field in ("title", "description"):
            self.out[field].append(data)


def fetch_page(url, timeout=15):
    """Read title/bullets/description from an Amazon product page. Returns (dict, error)."""
    req = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                      "Chrome/126.0 Safari/537.36",
        "Accept-Language": "en-US,en;q=0.9,zh-CN;q=0.8,ja;q=0.7,de;q=0.6"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            body = r.read(3_000_000).decode("utf-8", "replace")
            final_url = r.geturl()
    except Exception as e:  # noqa: BLE001 - any network failure is reported per row
        return None, f"商品页读取失败（{type(e).__name__}）"
    if "validateCaptcha" in body or "api-services-support@amazon.com" in body:
        return None, "商品页要求验证码，未读到内容"
    p = _PageParser()
    p.feed(body)
    title = " ".join("".join(p.out["title"]).split())
    if not title:
        return None, "商品页没有读到标题"
    return {"title": title, "bullets": p.out["bullets"], "final_url": final_url,
            "description": " ".join("".join(p.out["description"]).split())}, None


# ---------------------------------------------------------------- tables

def _col_index(ref):
    n = 0
    for ch in re.match(r"[A-Z]+", ref).group(0):
        n = n * 26 + ord(ch) - 64
    return n - 1


def _col_letter(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


def read_xlsx(path, sheet=None):
    """Rows (lists of strings) of one worksheet, standard library only. Returns (sheet_name, rows)."""
    with zipfile.ZipFile(path) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        sheets = [(s.get("name"), s.get("{%s}id" % NS["r"])) for s in wb.find("m:sheets", NS)]
        name, rid = next(((n, r) for n, r in sheets if n == sheet), sheets[0]) if sheet else sheets[0]
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        target = next(r.get("Target") for r in rels if r.get("Id") == rid)
        target = target.lstrip("/") if target.startswith("/") else "xl/" + target
        shared = []
        if "xl/sharedStrings.xml" in z.namelist():
            for si in ET.fromstring(z.read("xl/sharedStrings.xml")).findall("m:si", NS):
                shared.append("".join(t.text or "" for t in si.iter("{%s}t" % NS["m"])))
        rows = []
        for row in ET.fromstring(z.read(target)).iter("{%s}row" % NS["m"]):
            idx = int(row.get("r")) - 1 if row.get("r") else len(rows)
            while len(rows) < idx:
                rows.append([])
            vals = []
            for c in row.findall("m:c", NS):
                ci = _col_index(c.get("r")) if c.get("r") else len(vals)
                t, v = c.get("t"), c.find("m:v", NS)
                if t == "s" and v is not None:
                    val = shared[int(v.text)]
                elif t == "inlineStr":
                    val = "".join(x.text or "" for x in c.iter("{%s}t" % NS["m"]))
                else:
                    val = v.text if v is not None and v.text is not None else ""
                    if re.fullmatch(r"-?\d+\.0", val):
                        val = val[:-2]
                while len(vals) < ci:
                    vals.append("")
                vals.append(val)
            rows.append(vals)
    return name, rows


def _cell_xml(ref, value, style=0):
    if value is None or value == "":
        return ""
    s = f' s="{style}"' if style else ""
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return f'<c r="{ref}" t="inlineStr"{s}><is><t xml:space="preserve">{text}</t></is></c>'


def write_xlsx(path, sheets):
    """sheets: [(name, rows, widths)] -> minimal workbook, every cell as text (no formulas can sneak in)."""
    ct = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
          '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>']
    wb_sheets, wb_rels = [], []
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for i, (name, rows, widths) in enumerate(sheets, 1):
            ct.append(f'<Override PartName="/xl/worksheets/sheet{i}.xml" '
                      'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>')
            safe = re.sub(r"[\[\]:*?/\\]", "_", name)[:31].replace("&", "&amp;").replace('"', "&quot;")
            wb_sheets.append(f'<sheet name="{safe}" sheetId="{i}" r:id="rId{i}"/>')
            wb_rels.append(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                           f'relationships/worksheet" Target="worksheets/sheet{i}.xml"/>')
            cols = "".join(f'<col min="{j + 1}" max="{j + 1}" width="{w}" customWidth="1"/>'
                           for j, w in enumerate(widths or []) if w)
            body = []
            for r, row in enumerate(rows, 1):
                cells = "".join(_cell_xml(f"{_col_letter(c)}{r}", v, 1 if r == 1 else 2) for c, v in enumerate(row))
                body.append(f'<row r="{r}">{cells}</row>')
            z.writestr(f"xl/worksheets/sheet{i}.xml",
                       '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                       f'<worksheet xmlns="{NS["m"]}"><sheetViews><sheetView workbookViewId="0">'
                       '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>'
                       + (f"<cols>{cols}</cols>" if cols else "") + f'<sheetData>{"".join(body)}</sheetData></worksheet>')
        ct.append("</Types>")
        z.writestr("[Content_Types].xml", "".join(ct))
        z.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                   'officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr("xl/workbook.xml", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   f'<workbook xmlns="{NS["m"]}" xmlns:r="{NS["r"]}"><sheets>{"".join(wb_sheets)}</sheets></workbook>')
        n = len(sheets)
        z.writestr("xl/_rels/workbook.xml.rels", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   + "".join(wb_rels) + f'<Relationship Id="rId{n + 1}" Type="http://schemas.openxmlformats.org/'
                   'officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>')
        z.writestr("xl/styles.xml", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   f'<styleSheet xmlns="{NS["m"]}"><fonts count="2"><font><sz val="11"/></font>'
                   '<font><b/><sz val="11"/></font></fonts><fills count="2"><fill><patternFill patternType="none"/></fill>'
                   '<fill><patternFill patternType="gray125"/></fill></fills><borders count="1"><border/></borders>'
                   '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
                   '<cellXfs count="3"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
                   '<xf numFmtId="0" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1"/>'
                   '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyAlignment="1">'
                   '<alignment wrapText="1" vertical="top"/></xf></cellXfs>'
                   '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>')


def read_table(path, sheet=None):
    """(kind, sheet_name, rows). kind is "xlsx" or "csv"."""
    p = str(path).lower()
    if p.endswith((".xlsm", ".xls")):
        raise ValueError("不支持 .xlsm / .xls；请先在 Excel 中另存为 .xlsx（宏不会随格式转换保留）")
    if p.endswith(".xlsx"):
        runtime.require_excel()
        name, rows = read_xlsx(path, sheet)
        return "xlsx", name, rows
    if p.endswith((".csv", ".tsv", ".txt")):
        text = runtime.decode_text(pathlib.Path(path).read_bytes(), legacy_csv=True)
        dialect = "excel-tab" if p.endswith(".tsv") or ("\t" in text.split("\n", 1)[0] and "," not in text.split("\n", 1)[0]) else "excel"
        return "csv", None, [row for row in csv.reader(io.StringIO(text), dialect)]
    raise ValueError("只支持 .xlsx / .csv / .tsv 表格；.xls 请先另存为 .xlsx")


def _csv_safe(v):
    v = "" if v is None else str(v)
    return "'" + v if v[:1] in ("=", "+", "-", "@", "\t", "\r") else v


def write_table(src, kind, sheet_name, header_row, rows, new_header, new_cols, extra_sheets, out_path):
    """Append new columns to the source table and save to out_path. Returns a note on how it was written."""
    target = pathlib.Path(out_path)
    allowed = (".csv", ".tsv") if kind == "csv" else (".xlsx",)
    if target.suffix.lower() not in allowed:
        raise ValueError("输出扩展名必须为 " + " / ".join(allowed))
    if target.resolve() == pathlib.Path(src).resolve():
        raise ValueError("输出文件不能覆盖原表")
    if kind != "csv":
        runtime.require_excel()
    # Close before reopening: Windows does not permit reopening NamedTemporaryFile's handle.
    with tempfile.NamedTemporaryFile(dir=target.parent, prefix=".csc-", suffix=target.suffix, delete=False) as fh:
        temporary = pathlib.Path(fh.name)
    try:
        note = _write_table(src, kind, sheet_name, header_row, rows, new_header, new_cols, extra_sheets, temporary)
        temporary.replace(target)
        return note
    finally:
        temporary.unlink(missing_ok=True)


def _write_table(src, kind, sheet_name, header_row, rows, new_header, new_cols, extra_sheets, out_path):
    if kind == "csv":
        with open(out_path, "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh, delimiter="\t" if pathlib.Path(out_path).suffix.lower() == ".tsv" else ",")
            width = max(len(r) for r in rows) if rows else 0
            for i, r in enumerate(rows):
                pad = list(r) + [""] * (width - len(r))
                if i < header_row:
                    w.writerow(pad)
                elif i == header_row:
                    w.writerow(pad + new_header)
                else:
                    w.writerow(pad + [_csv_safe(v) for v in new_cols.get(i, [""] * len(new_header))])
        return "csv"
    openpyxl = runtime.require_excel()
    from openpyxl.styles import Alignment, Font
    if kind == "xlsx":
        wb = openpyxl.load_workbook(src)
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]
        start = max(ws.max_column, max((len(r) for r in rows), default=0)) + 1
        for j, h in enumerate(new_header):
            c = ws.cell(row=header_row + 1, column=start + j, value=h)
            c.font = Font(bold=True)
            ws.column_dimensions[_col_letter(start + j - 1)].width = (14, 16, 70)[min(j, 2)]
        for i, vals in new_cols.items():
            for j, v in enumerate(vals):
                c = ws.cell(row=i + 1, column=start + j)
                c.value = v
                c.data_type = "s"
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for name, srows, widths in extra_sheets:
            if name in wb.sheetnames:
                del wb[name]
            es = wb.create_sheet(name)
            for r in srows:
                es.append([str(x) for x in r])
            for j, w in enumerate(widths):
                es.column_dimensions[_col_letter(j)].width = w
        try:
            wb.save(out_path)
        finally:
            wb.close()
        return "xlsx（保留原表格式）"
    raise ValueError("未知表格类型：" + str(kind))


HEADER_KEYS = {
    "title": ("商品标题", "产品标题", "标题", "商品名称", "产品名称", "品名", "title", "product name", "item name"),
    "url": ("商品链接", "产品链接", "链接", "网址", "url", "link"),
    "asin": ("asin",),
    "bullets": ("亮点", "五点", "卖点", "要点", "bullet", "highlight"),
    "description": ("描述", "description"),
    "site": ("站点", "site", "marketplace", "国家"),
}


def detect_columns(header, overrides=None):
    """Map field -> list of column indexes from header names (overrides: field -> header name or column letter)."""
    cols = {}
    norm = [str(h or "").strip().lower() for h in header]
    for field, keys in HEADER_KEYS.items():
        o = (overrides or {}).get(field)
        if o:
            if re.fullmatch(r"[A-Za-z]{1,3}", o) and o.lower() not in norm:
                cols[field] = [_col_index(o.upper())]
            elif o.strip().lower() in norm:
                cols[field] = [norm.index(o.strip().lower())]
            else:
                raise ValueError(f"找不到列：{o}")
            continue
        hits = []
        for i, h in enumerate(norm):
            if not h:
                continue
            if field == "site" and any(k in h for k in ("卖家", "供应商", "产地", "来源", "seller", "supplier", "origin", "source")):
                continue
            if field == "description" and any(k in h for k in HEADER_KEYS["bullets"]):
                continue
            if field == "title" and any(k in h for k in ("链接", "url", "asin")):
                continue
            if any(k in h for k in keys) or (field == "title" and h in ("name", "名称")):
                hits.append(i)
        if hits:
            # prefer the product's own column over source/list/store columns ("来源店铺URL列表" vs "商品URL")
            noise = ("来源", "列表", "店铺", "搜索", "source", "store", "list")
            hits.sort(key=lambda i: (any(n in norm[i] for n in noise), i))
            cols[field] = sorted(hits) if field == "bullets" else hits[:1]
    return cols


def find_header(rows, limit=10):
    """Index of the header row: the first of the first rows that names a title or URL column."""
    for i, r in enumerate(rows[:limit]):
        c = detect_columns(r)
        if "title" in c or "url" in c:
            return i
    return 0


def politeness(last, gap=1.5):
    wait = gap - (time.time() - last)
    if wait > 0:
        time.sleep(wait)
    return time.time()
