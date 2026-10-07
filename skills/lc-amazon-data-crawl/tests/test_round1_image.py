"""Round-1 fixes for the image-competitor crawler (paid-call safety, D2
quarantine, Find Similar fallbacks, budgets, dedupe, exit contract)."""

from __future__ import annotations

import io
import json
import os
from contextlib import redirect_stdout
from pathlib import Path
from types import SimpleNamespace
import sys
import tempfile
import unittest
from unittest.mock import MagicMock, patch


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = SKILL_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import amazon_image_competitor_crawler as image
import run_outcome


class FakeResponse:
    def __init__(self, status_code: int, payload: object | None = None, headers: dict | None = None) -> None:
        self.status_code = status_code
        self._payload = payload
        self.text = ""
        self.headers = headers or {}

    def json(self) -> object:
        return self._payload


def mini_content_response(content: str) -> FakeResponse:
    return FakeResponse(200, {"choices": [{"message": {"content": content}}]})


def candidate(index: int, *, source_asin: str = "B0SOURCE00") -> dict:
    return {
        "asin": f"B{index:09d}",
        "source_asin": source_asin,
        "candidate_image_url": f"https://m.media-amazon.com/images/I/{index}.jpg",
        "title": f"candidate {index}",
        "rank": str(index),
    }


def cascade_runtime(**overrides: object) -> SimpleNamespace:
    raw: dict = {
        "match_mode": "cascade",
        "include_source_as_competitor": False,
        "prescreen_min_similarity": 0.7,
        "prescreen_max_matches": 10,
        "mini_batch_size": 6,
        "mini_retry_attempts": 3,
        "mini_retry_backoff_seconds": 0,
        "mini_api_key": "mini-key",
        "mini_model": "mini-model",
        "mini_base_url": "https://ark.example/api",
        "mini_api_path": "chat/completions",
        "vision_timeout": 5,
        "provider_metrics": {},
        "embedding_provider": "doubao",
        "embedding_base_url": "https://ark.example/api",
        "embedding_api_path": "embeddings/multimodal",
        "embedding_model": "emb-model",
        "embedding_encoding_format": "float",
    }
    raw.update(overrides)
    return SimpleNamespace(**raw)


class IsolatedHomeTestCase(unittest.TestCase):
    """Keep the global safety lock/traffic files away from ~ and other runs."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        home = self.root / "home"
        home.mkdir()
        env = patch.dict(os.environ, {"HOME": str(home)})
        env.start()
        self.addCleanup(env.stop)
        pacing = patch("safety_control.LocalSafetyController._wait_until")
        pacing.start()
        self.addCleanup(pacing.stop)
        image.EMBEDDING_CACHE.clear()
        self.addCleanup(image.EMBEDDING_CACHE.clear)
        self.source_image = self.root / "source.jpg"
        self.source_image.write_bytes(b"\xff\xd8\xff source")

    def tearDown(self) -> None:
        self.temp_dir.cleanup()


# ---------------------------------------------------------------------------
# Item 1 (L11) + item 3: paid results are checkpointed per source.
# ---------------------------------------------------------------------------


class PaidCallCheckpointTests(IsolatedHomeTestCase):
    def store(self) -> image.SourceProgressStore:
        return image.SourceProgressStore(
            self.root / "job" / "source_progress" / "00000002-abc",
            {"source_id": "B0SOURCE00#row-2", "provider_sha256": "p", "crawl_plan_sha256": "c"},
        )

    def test_findings_repro_restart_makes_zero_duplicate_embedding_calls(self) -> None:
        """Repro from findings (image_mini_malformed.py) with a deterministic
        malformed Mini answer: run 1 pays N embeddings, the restart pays 0."""

        calls = {"emb": 0, "mini": 0}
        payloads: list[str] = []

        def fake_emb(_runtime: object, _ref: str) -> list[float]:
            calls["emb"] += 1
            return [1.0, 0.0]

        def fake_post(url, headers=None, json=None, timeout=None, allow_redirects=None):  # noqa: A002
            calls["mini"] += 1
            payloads.append(repr(json))
            return mini_content_response("完全不是 JSON")

        records = [candidate(i) for i in range(8)]
        with (
            patch.object(image, "call_multimodal_embedding", side_effect=fake_emb),
            patch.object(image.requests, "post", side_effect=fake_post),
            patch.object(image.time, "sleep"),
        ):
            for run in (1, 2):
                image.EMBEDDING_CACHE.clear()  # new process
                runtime = cascade_runtime()
                runtime._source_progress = self.store()
                with self.assertRaises(image.MiniProviderError):
                    image.run_cascade_match(runtime, self.source_image, "", records)
                if run == 1:
                    first_run_emb = calls["emb"]
                    first_run_mini = calls["mini"]
        self.assertEqual(first_run_emb, 9)  # source + 8 candidates
        self.assertEqual(calls["emb"], first_run_emb, "restart must not re-pay any embedding")
        # Malformed temperature-0 output is sent at most twice per attempt and
        # the resend is not identical (it carries a JSON repair instruction).
        self.assertEqual(first_run_mini, 2)
        self.assertEqual(calls["mini"], 4)
        self.assertNotEqual(payloads[0], payloads[1])
        self.assertIn(image.MINI_JSON_REPAIR_TEXT, payloads[1])

    def test_fenced_mini_json_from_findings_repro_is_accepted(self) -> None:
        records = [candidate(1)]
        fenced = "```json\n" + json.dumps(
            {"matches": [{"asin": records[0]["asin"], "is_same_product": True, "confidence": 0.9, "reason": "同款"}]},
            ensure_ascii=False,
        ) + "\n```"
        with (
            patch.object(image, "call_multimodal_embedding", return_value=[1.0, 0.0]),
            patch.object(image.requests, "post", return_value=mini_content_response(fenced)) as post,
        ):
            result = image.run_cascade_match(cascade_runtime(), self.source_image, "", records)
        self.assertEqual(post.call_count, 1)
        self.assertEqual(result.same_product_count, 1)

    def test_completed_mini_batch_is_never_resent_after_a_later_batch_fails(self) -> None:
        records = [candidate(i) for i in range(1, 9)]  # 8 matches -> batches of 6 + 2
        sent_batches: list[list[str]] = []
        fail_second = {"value": True}

        def fake_mini(_runtime, _path, batch):
            asins = [row["asin"] for row in batch]
            sent_batches.append(asins)
            if len(batch) == 2 and fail_second["value"]:
                raise image.MiniProviderError("HTTP 503")
            return [
                {"asin": asin, "is_same_product": True, "confidence": 0.9, "reason": "同款"}
                for asin in asins
            ]

        with (
            patch.object(image, "call_multimodal_embedding", return_value=[1.0, 0.0]),
            patch.object(image, "call_doubao_mini_verifier", side_effect=fake_mini),
        ):
            runtime = cascade_runtime()
            runtime._source_progress = self.store()
            with self.assertRaises(image.MiniProviderError):
                image.run_cascade_match(runtime, self.source_image, "", records)
            fail_second["value"] = False
            image.EMBEDDING_CACHE.clear()
            runtime = cascade_runtime()
            runtime._source_progress = self.store()
            result = image.run_cascade_match(runtime, self.source_image, "", records)
        self.assertEqual([len(batch) for batch in sent_batches], [6, 2, 2])
        self.assertEqual(result.same_product_count, 8)
        self.assertEqual(runtime.provider_metrics.get("mini_reused_from_checkpoint"), 1)

    def test_store_is_discarded_when_provider_fingerprint_changes(self) -> None:
        store = self.store()
        store.put("emb", "k", {"vector": [1.0]})
        same = self.store()
        self.assertIsNotNone(same.get("emb", "k"))
        changed = image.SourceProgressStore(
            store.directory,
            {"source_id": "B0SOURCE00#row-2", "provider_sha256": "OTHER", "crawl_plan_sha256": "c"},
        )
        self.assertIsNone(changed.get("emb", "k"))


# ---------------------------------------------------------------------------
# Item 2: candidate exclusion vs. source failure.
# ---------------------------------------------------------------------------


class CandidateExclusionTests(IsolatedHomeTestCase):
    def test_unembeddable_candidate_is_excluded_and_not_repaid(self) -> None:
        # 2 of 5 unscorable (<= 50 %): excluded with a note.  More than half
        # unscorable fails the source instead (round 1b, P1-6).
        records = [
            candidate(1),
            candidate(2),
            {**candidate(3), "candidate_image_url": ""},
            candidate(4),
            candidate(5),
        ]
        bad_url = records[1]["candidate_image_url"]
        calls: list[str] = []

        def fake_emb(_runtime, ref):
            calls.append(ref)
            if ref == bad_url:
                raise image.EmbeddingInputRejected("HTTP 400")
            return [1.0, 0.0] if not ref.startswith("https://m.media") else [0.0, 1.0]

        store_dir = self.root / "progress"
        identity = {"source_id": "s", "provider_sha256": "p", "crawl_plan_sha256": "c"}
        with (
            patch.object(image, "call_multimodal_embedding", side_effect=fake_emb),
            patch.object(image, "image_url_to_data_url", side_effect=image.requests.ConnectionError("down")),
        ):
            runtime = cascade_runtime()
            runtime._source_progress = image.SourceProgressStore(store_dir, identity)
            result = image.run_cascade_match(runtime, self.source_image, "", records)
            first_calls = len(calls)
            image.EMBEDDING_CACHE.clear()
            runtime = cascade_runtime()
            runtime._source_progress = image.SourceProgressStore(store_dir, identity)
            image.run_cascade_match(runtime, self.source_image, "", records)
        self.assertEqual(result.processing_status, "verified_zero")
        self.assertEqual(result.decisions[records[1]["asin"]]["prescreen_status"], "unscorable")
        self.assertEqual(result.decisions[records[2]["asin"]]["prescreen_status"], "unscorable")
        self.assertIn("无法识别", result.match_reason)
        self.assertEqual(len(calls), first_calls, "excluded image must not be re-sent after restart")

    def test_transient_provider_failure_is_a_source_failure_not_an_exclusion(self) -> None:
        with patch.object(
            image,
            "call_multimodal_embedding_cached",
            side_effect=image.EmbeddingProviderError("HTTP 503 after retries"),
        ), patch.object(image, "image_url_to_data_url") as fallback:
            with self.assertRaises(image.EmbeddingProviderError):
                image.resolve_candidate_embedding_for_cascade(cascade_runtime(), candidate(1))
        fallback.assert_not_called()


# ---------------------------------------------------------------------------
# Item 3: 429 backoff, Retry-After.
# ---------------------------------------------------------------------------


class ProviderBackoffTests(IsolatedHomeTestCase):
    def test_429_waits_at_least_ten_seconds_growing_and_honours_retry_after(self) -> None:
        r429 = FakeResponse(429)
        self.assertEqual(image.provider_retry_wait_seconds(r429, 0, 1.0), 10.0)
        self.assertEqual(image.provider_retry_wait_seconds(r429, 1, 1.0), 20.0)
        self.assertEqual(
            image.provider_retry_wait_seconds(FakeResponse(429, headers={"Retry-After": "45"}), 0, 1.0),
            45.0,
        )
        self.assertEqual(image.provider_retry_wait_seconds(FakeResponse(503), 1, 1.0), 2.0)

    def test_mini_429_uses_retry_after_and_heartbeats_long_waits(self) -> None:
        rows = [candidate(1)]
        ok = mini_content_response(
            json.dumps({"matches": [{"asin": rows[0]["asin"], "is_same_product": False, "confidence": 0.1, "reason": "不同"}]})
        )
        runtime = cascade_runtime(mini_retry_backoff_seconds=1)
        runtime.safety = MagicMock(spec=image.LocalSafetyController)
        with (
            patch.object(image.requests, "post", side_effect=[FakeResponse(429, headers={"Retry-After": "40"}), ok]),
            patch.object(image.time, "sleep") as sleep,
            redirect_stdout(io.StringIO()),
        ):
            image.call_doubao_mini_verifier(runtime, self.source_image, rows)
        self.assertEqual(sum(call.args[0] for call in sleep.call_args_list), 40)
        self.assertLessEqual(max(call.args[0] for call in sleep.call_args_list), 30)
        self.assertGreaterEqual(runtime.safety.heartbeat.call_count, 2)


# ---------------------------------------------------------------------------
# Items 4-6: Find Similar paths, budgets, element waits.
# ---------------------------------------------------------------------------


class FindSimilarTests(IsolatedHomeTestCase):
    def runtime(self) -> SimpleNamespace:
        return SimpleNamespace(
            search_strategy="sellersprite_find_similar_first",
            marketplace_domain="amazon.com",
            page_timeout=90,
            find_similar_timeout=20,
        )

    def test_page_not_found_in_find_similar_takes_canonical_check_then_source_unavailable(self) -> None:
        class Driver:
            title = "Page Not Found"
            current_url = "about:blank"
            window_handles: list = []

        opened: list[str] = []

        def fake_open(_driver, url, _runtime, state=None):
            opened.append(url)
            raise image.TransientAmazonPageUnavailable("dog", reason="amazon_dog_error", url=url)

        current = {
            "source_asin": "B000000001",
            "source_product_url": "https://www.amazon.com/Some-Title/dp/B000000001?ref=sr_1",
        }
        with (
            patch.object(image, "open_image_amazon_page", side_effect=fake_open),
            patch.object(image, "upload_image_to_lens") as upload,
        ):
            with self.assertRaises(image.SourceProductUnavailable):
                image.run_image_search(Driver(), self.runtime(), current, Path("x.jpg"))
        self.assertEqual(opened, [current["source_product_url"], "https://www.amazon.com/dp/B000000001"])
        upload.assert_not_called()

    def test_transient_product_page_in_find_similar_falls_back_to_lens_upload(self) -> None:
        class Driver:
            title = "Amazon.com"
            current_url = "about:blank"
            window_handles: list = []

        runtime = self.runtime()
        current = {"source_asin": "B000000001", "source_product_url": "https://www.amazon.com/dp/B000000001"}
        with (
            patch.object(
                image,
                "open_image_amazon_page",
                # Round 1b: only a page without the expected content falls back;
                # 5xx / navigation errors go to the stage retry (test_round1b_image).
                side_effect=image.TransientAmazonPageUnavailable(
                    "no main image", reason="expected_content_missing"
                ),
            ),
            patch.object(image, "upload_image_to_lens") as upload,
        ):
            method = image.run_image_search(Driver(), runtime, current, Path("x.jpg"))
        self.assertEqual(method, "amazon_upload")
        upload.assert_called_once()
        self.assertTrue(current["find_similar_failed"])

    def test_missing_control_spends_no_counted_action(self) -> None:
        class Driver:
            current_url = "https://www.amazon.com/dp/B000000001"
            title = "Amazon.com: thing"
            window_handles = ["w"]

            def execute_script(self, script, *_args):
                if script == image.FIND_SIMILAR_CONTROL_SCRIPT:
                    return False
                return "https://m.media-amazon.com/images/I/x.jpg"

        class ImmediateWait:
            def __init__(self, *_a, **_k):
                pass

            def until(self, predicate):
                return True

        runtime = self.runtime()
        runtime.safety = MagicMock(spec=image.LocalSafetyController)
        current = {"source_asin": "B000000001", "source_product_url": "https://www.amazon.com/dp/B000000001"}
        with (
            patch.object(image, "WebDriverWait", ImmediateWait),
            patch.object(image, "FIND_SIMILAR_CONTROL_WAIT_SECONDS", 0),
            patch.object(image.time, "sleep"),
            patch.object(image, "upload_image_to_lens") as upload,
            redirect_stdout(io.StringIO()),
        ):
            method = image.run_image_search(Driver(), runtime, current, Path("x.jpg"))
        self.assertEqual(method, "amazon_upload")
        runtime.safety.before_remote_action.assert_not_called()
        upload.assert_called_once()

    def test_find_similar_is_sticky_per_source_and_disabled_after_three_failures(self) -> None:
        runtime = self.runtime()
        with (
            patch.object(image, "trigger_sellersprite_find_similar", return_value=False) as trigger,
            patch.object(image, "upload_image_to_lens"),
            redirect_stdout(io.StringIO()),
        ):
            sticky = {"source_asin": "B000000009"}
            image.run_image_search(MagicMock(), runtime, sticky, Path("x.jpg"))
            image.run_image_search(MagicMock(), runtime, sticky, Path("x.jpg"))
            self.assertEqual(trigger.call_count, 1, "retry of the same source must not repeat Find Similar")
            for index in range(2, 6):
                image.run_image_search(MagicMock(), runtime, {"source_asin": f"B00000000{index}"}, Path("x.jpg"))
        # sources 1..3 tried FS; after 3 consecutive failures the rest skip it.
        self.assertEqual(trigger.call_count, 3)

    def test_source_budget_excludes_throttle_and_stops_long_retry_waits(self) -> None:
        clock = {"now": 0.0}
        budget = image.SourceBudget(100, clock=lambda: clock["now"])
        clock["now"] = 150.0
        budget.exclude(60.0)
        budget.check("ok")  # 90 s of budgeted time
        clock["now"] = 170.0
        with self.assertRaises(image.SourceTimeBudgetExceeded):
            budget.check("over")

        retry_clock = {"now": 1000.0}

        def waiter(seconds: float) -> None:
            retry_clock["now"] += seconds
            clock["now"] += seconds

        clock["now"] = 0.0
        runtime = SimpleNamespace(
            marketplace_domain="amazon.com",
            amazon_page_unavailable_retry_schedule_seconds=((300.0, 300.0),),
            _amazon_page_retry_clock=lambda: retry_clock["now"],
            _amazon_page_retry_waiter=waiter,
            _amazon_page_retry_rng=SimpleNamespace(uniform=lambda a, b: a),
            _source_budget=image.SourceBudget(100, clock=lambda: clock["now"]),
        )
        state = MagicMock()
        state.load_amazon_page_retry.return_value = None
        attempts: list[int] = []

        def failing(attempt):
            attempts.append(1)
            raise image.TransientAmazonPageUnavailable("blank", reason="blank_page")

        with self.assertRaises(image.SourceTimeBudgetExceeded), redirect_stdout(io.StringIO()):
            image.run_image_page_stage_with_recovery(
                runtime, state, {"source_id": "S"}, "lens_results", "https://www.amazon.com/x", lambda: None, failing
            )
        self.assertEqual(len(attempts), 1)
        self.assertLessEqual(clock["now"], 100.0)

    def test_counted_action_is_not_spent_when_budget_is_gone(self) -> None:
        runtime = SimpleNamespace(
            safety=MagicMock(spec=image.LocalSafetyController),
            _source_budget=image.SourceBudget(10, clock=lambda: 50.0),
        )
        runtime._source_budget.started = 0.0
        with self.assertRaises(image.SourceTimeBudgetExceeded):
            image.counted_remote_action(runtime, "导航")
        runtime.safety.before_remote_action.assert_not_called()

    def test_loaded_page_element_waits_are_at_most_thirty_seconds(self) -> None:
        timeouts: list[float] = []

        class RecordingWait:
            def __init__(self, _driver, timeout, *_a, **_k):
                timeouts.append(float(timeout))

            def until(self, predicate):
                return True

        driver = MagicMock()
        driver.title = "Amazon.com: item"
        driver.current_url = "https://www.amazon.com/dp/B000000001"
        driver.execute_script.return_value = "https://m.media-amazon.com/images/I/x.jpg"
        runtime = SimpleNamespace(page_timeout=90)
        with (
            patch.object(image, "open_image_amazon_page"),
            patch.object(image, "WebDriverWait", RecordingWait),
        ):
            image.load_source_product_main_image(driver, runtime, "https://www.amazon.com/dp/B000000001", None)
            image.open_find_similar_product_page(
                driver,
                SimpleNamespace(page_timeout=90, marketplace_domain="amazon.com"),
                {"source_asin": "B000000001"},
                "https://www.amazon.com/dp/B000000001",
            )
        self.assertTrue(timeouts)
        self.assertLessEqual(max(timeouts), 30.0)


# ---------------------------------------------------------------------------
# Item 7: progress heartbeat; item 10: retry heartbeat shows the real max.
# ---------------------------------------------------------------------------


class ProgressOutputTests(IsolatedHomeTestCase):
    def test_sequential_embedding_calls_heartbeat_every_ten(self) -> None:
        runtime = cascade_runtime()
        runtime.safety = MagicMock(spec=image.LocalSafetyController)
        records = [candidate(i) for i in range(1, 13)]
        out = io.StringIO()
        with patch.object(image, "call_multimodal_embedding_cached", side_effect=[[1.0, 0.0]] + [[0.0, 1.0]] * 12), redirect_stdout(out):
            image.run_cascade_match(runtime, self.source_image, "", records)
        phases = [call.args[0] for call in runtime.safety.heartbeat.call_args_list]
        self.assertGreaterEqual(phases.count("provider_calls"), 2)
        self.assertIn("10/12", out.getvalue())

    def test_retry_heartbeat_prints_real_max_attempts(self) -> None:
        clock = {"now": 1000.0}

        def waiter(seconds: float) -> None:
            clock["now"] += seconds

        runtime = SimpleNamespace(
            marketplace_domain="amazon.com",
            amazon_page_unavailable_retry_schedule_seconds=((60.0, 60.0),),
            _amazon_page_retry_clock=lambda: clock["now"],
            _amazon_page_retry_waiter=waiter,
            _amazon_page_retry_rng=SimpleNamespace(uniform=lambda a, b: a),
        )
        state = MagicMock()
        state.load_amazon_page_retry.return_value = None
        attempts: list[int] = []

        def flaky(attempt):
            attempts.append(1)
            if len(attempts) == 1:
                raise image.TransientAmazonPageUnavailable("blank", reason="blank_page")
            return "ok"

        out = io.StringIO()
        with redirect_stdout(out):
            image.run_image_page_stage_with_recovery(
                runtime, state, {"source_id": "S"}, "lens_results", "https://www.amazon.com/x", lambda: None, flaky
            )
        self.assertIn("下一次=2/2", out.getvalue())
        self.assertNotIn("/5", out.getvalue())


class StateStoreTests(unittest.TestCase):
    def test_deferred_sources_are_restored_at_the_end_of_the_queue(self) -> None:
        store = image.ImageCompetitorStateStore.__new__(image.ImageCompetitorStateStore)
        store.data = {"queue": [{"source_id": "X"}], "deferred_sources": [{"source": {"source_id": "A"}, "reason": "r"}]}
        store.flush = lambda: None  # type: ignore[assignment]
        store.restore_deferred()
        self.assertEqual([item["source_id"] for item in store.data["queue"]], ["X", "A"])


# ---------------------------------------------------------------------------
# Full-loop tests: D2 quarantine, deferral, dedupe, download fallback,
# workbook on early exit, run_summary/exit codes.
# ---------------------------------------------------------------------------


class LoopTests(IsolatedHomeTestCase):
    def setUp(self) -> None:
        super().setUp()
        self.embedding_config = self.root / "embedding.json"
        self.embedding_config.write_text(json.dumps({"api_key": "emb-secret"}), encoding="utf-8")
        self.mini_config = self.root / "mini.json"
        self.mini_config.write_text(json.dumps({"api_key": "mini-secret"}), encoding="utf-8")

    def write_products(self, rows: list[str], header: str = "ASIN,商品URL,主图URL,本地图片路径,备注") -> Path:
        path = self.root / "products.csv"
        path.write_text(header + "\n" + "\n".join(rows) + "\n", encoding="utf-8")
        return path

    def raw_config(self, products: Path, **overrides: object) -> dict:
        raw: dict = {
            "job_id": "round1-image",
            "outputs_root": str(self.root / "outputs"),
            "products_file": str(products),
            "marketplace": "美国站",
            "result_mode": "count_only",
            "match_mode": "cascade",
            "doubao_embedding_config_file": str(self.embedding_config),
            "doubao_mini_config_file": str(self.mini_config),
            "mini_retry_backoff_seconds": 0,
            "embedding_retry_backoff_seconds": 0,
            "search_strategy": "amazon_upload",
            "browser_backend": "cdp",
            "browser_mode": "reuse",
            "extension_path": "auto",
            "save_debug_snapshots": False,
        }
        raw.update(overrides)
        return raw

    def runtime(self, products: Path, **overrides: object) -> image.ImageCompetitorRuntimeConfig:
        runtime = image.build_image_runtime_config(self.raw_config(products, **overrides), no_resume=False)
        runtime.amazon_page_unavailable_retry_schedule_seconds = ((0, 0),)
        return runtime

    def job_dir(self, runtime: image.ImageCompetitorRuntimeConfig) -> Path:
        return runtime.outputs_root / runtime.job_id

    def loop_patches(self, merge_side_effect=None):
        driver = MagicMock()
        driver.current_url = "https://www.amazon.com/products?searchtype=flow"
        merge = (lambda _d, _r, current, _s: [
            {**candidate(1, source_asin=current["source_asin"]), "source_id": current["source_id"]},
            {**candidate(2, source_asin=current["source_asin"]), "source_id": current["source_id"]},
            {**candidate(3, source_asin=current["source_asin"]), "source_id": current["source_id"]},
        ])
        return [
            patch.object(image, "start_driver", return_value=driver),
            patch.object(image, "resolve_source_image", return_value=self.source_image),
            patch.object(image, "run_image_search", return_value="amazon_upload"),
            patch.object(image, "wait_for_lens_results", return_value="results"),
            patch.object(image, "merge_lens_product_data", side_effect=merge_side_effect or merge),
        ]

    def run_crawl(self, runtime, extra_patches=()):
        patches = self.loop_patches() + list(extra_patches)
        started = [p.start() for p in patches]
        try:
            with redirect_stdout(io.StringIO()):
                return image.run_image_competitor_crawl(runtime, dry_run=False), started
        finally:
            for p in reversed(patches):
                p.stop()

    def test_provider_outage_retry_later_never_quarantines_and_never_repays(self) -> None:
        # Round 1b (P1-3): a provider 503 is an environment failure; two
        # failed runs no longer skip the source, and paid calls stay saved.
        products = self.write_products(["B012345678,,,,x"])
        runtime = self.runtime(products)
        calls: list[str] = []
        outage = {"on": True}

        def emb(_runtime, ref):
            calls.append(ref)
            if ref.endswith("/3.jpg") and outage["on"]:
                raise image.EmbeddingProviderError("豆包视觉向量服务重试后仍不可用（HTTP 503）。")
            return [1.0, 0.0] if ref.startswith("data:") else [0.0, 1.0]

        with self.assertRaises(run_outcome.CrawlStop) as first:
            self.run_crawl(runtime, [patch.object(image, "call_multimodal_embedding", side_effect=emb)])
        self.assertEqual(first.exception.exit_code, run_outcome.EXIT_RETRY_LATER)
        job_dir = self.job_dir(runtime)
        state = image.load_json(job_dir / "state.json")
        self.assertIsNotNone(state["current"])
        source_id = state["current"]["source_id"]
        self.assertEqual(state.get("item_failure_cycles") or {}, {})
        self.assertEqual(state["item_failure_meta"][source_id]["environment_runs"], 1)
        self.assertTrue(state["last_failure_environment"])
        self.assertEqual(len(calls), 4)  # source, 1, 2, 3(failed)

        image.EMBEDDING_CACHE.clear()
        with self.assertRaises(run_outcome.CrawlStop) as second:
            self.run_crawl(self.runtime(products), [patch.object(image, "call_multimodal_embedding", side_effect=emb)])
        self.assertEqual(second.exception.exit_code, run_outcome.EXIT_RETRY_LATER)
        # Restart: only the failing call is re-sent; the source is NOT skipped.
        self.assertEqual(len(calls), 5)
        state = image.load_json(job_dir / "state.json")
        self.assertIsNotNone(state["current"])
        self.assertFalse(state.get("skipped_items"))
        self.assertEqual(state["item_failure_meta"][source_id]["environment_runs"], 2)

        outage["on"] = False
        image.EMBEDDING_CACHE.clear()
        result, started = self.run_crawl(
            self.runtime(products), [patch.object(image, "call_multimodal_embedding", side_effect=emb)]
        )
        self.assertEqual(result, 0)
        self.assertEqual(len(calls), 6)
        started[2].assert_not_called()  # run_image_search: no navigation repeated
        state = image.load_json(job_dir / "state.json")
        self.assertIsNone(state["current"])
        self.assertFalse(state.get("skipped_items"))
        counts = image.read_jsonl(job_dir / "counts.jsonl")
        self.assertEqual(counts[0]["processing_status"], "verified_zero")
        self.assertEqual(image._skipped_count(job_dir), 0)
        self.assertFalse((job_dir / image.SOURCE_PROGRESS_DIRNAME).exists() and any((job_dir / image.SOURCE_PROGRESS_DIRNAME).iterdir()))

    def test_fatal_provider_error_is_needs_human_with_config_next_action(self) -> None:
        mapped = image.classify_crawl_exception(image.FatalMiniProviderError("HTTP 401"))
        self.assertIsInstance(mapped, run_outcome.CrawlStop)
        self.assertEqual(mapped.exit_code, run_outcome.EXIT_NEEDS_HUMAN)
        self.assertIn("config/doubao_", mapped.next_action)
        self.assertIn("api_key", mapped.next_action)
        self.assertEqual(
            image.classify_crawl_exception(image.VerificationUnconfirmedError("x")).exit_code,
            run_outcome.EXIT_NEEDS_HUMAN,
        )
        self.assertEqual(
            image.classify_crawl_exception(image.UserFacingError("bad config")).exit_code,
            run_outcome.EXIT_CONFIG_ERROR,
        )

    def test_unattended_defers_failed_source_to_end_and_ends_with_retry_later(self) -> None:
        products = self.write_products(["B0AAAAAAAA,,,,x", "B0BBBBBBBB,,,,x"])
        runtime = self.runtime(products, operation_mode="unattended")
        runtime.amazon_page_unavailable_retry_schedule_seconds = ((0, 0),)

        def emb(_runtime, ref):
            if ref.endswith("/1.jpg") and "fail" in calls_mode:
                raise image.EmbeddingProviderError("HTTP 429 after retries")
            return [1.0, 0.0] if ref.startswith("data:") else [0.0, 1.0]

        calls_mode = {"fail"}
        seen_sources: list[str] = []

        def merge(_d, _r, current, _s):
            seen_sources.append(current["source_asin"])
            rows = [{**candidate(1, source_asin=current["source_asin"]), "source_id": current["source_id"]}]
            if current["source_asin"] == "B0BBBBBBBB":
                rows[0]["candidate_image_url"] = "https://m.media-amazon.com/images/I/ok.jpg"
            return rows

        patches = self.loop_patches(merge) + [patch.object(image, "call_multimodal_embedding", side_effect=emb)]
        for p in patches:
            p.start()
        try:
            with redirect_stdout(io.StringIO()), self.assertRaises(run_outcome.CrawlStop) as stop:
                image.run_image_competitor_crawl(runtime, dry_run=False)
            self.assertEqual(stop.exception.exit_code, run_outcome.EXIT_RETRY_LATER)
            job_dir = self.job_dir(runtime)
            state = image.load_json(job_dir / "state.json")
            self.assertEqual(len(state["deferred_sources"]), 1)
            self.assertEqual(seen_sources, ["B0AAAAAAAA", "B0BBBBBBBB"])
            # Workbook exists although the run did not finish cleanly (item 9).
            self.assertTrue(list(job_dir.glob("*.xlsx")))
            # Round 1b (P1-3): the provider 429 is an environment failure, so a
            # second failed run defers the source again instead of skipping it.
            with redirect_stdout(io.StringIO()), self.assertRaises(run_outcome.CrawlStop) as again:
                image.run_image_competitor_crawl(
                    self.runtime(products, operation_mode="unattended"), dry_run=False
                )
            self.assertEqual(again.exception.exit_code, run_outcome.EXIT_RETRY_LATER)
            self.assertFalse(image.load_json(job_dir / "state.json").get("skipped_items"))
            calls_mode.clear()
            with redirect_stdout(io.StringIO()):
                self.assertEqual(image.run_image_competitor_crawl(
                    self.runtime(products, operation_mode="unattended"), dry_run=False
                ), 0)
        finally:
            for p in reversed(patches):
                p.stop()
        state = image.load_json(job_dir / "state.json")
        self.assertFalse(state.get("skipped_items"))
        self.assertEqual(len(image.read_jsonl(job_dir / "counts.jsonl")), 2)

    def test_duplicate_input_asins_are_processed_once_in_count_only(self) -> None:
        products = self.write_products(["B012345678,,,,a", "B087654321,,,,b", "B012345678,,,,dup"])
        runtime = self.runtime(products)

        def emb(_runtime, ref):
            return [1.0, 0.0] if ref.startswith("data:") else [0.0, 1.0]

        _result, started = self.run_crawl(runtime, [patch.object(image, "call_multimodal_embedding", side_effect=emb)])
        self.assertEqual(started[1].call_count, 2)  # resolve_source_image
        counts = {row["input_row"]: row for row in image.read_jsonl(self.job_dir(runtime) / "counts.jsonl")}
        self.assertEqual(sorted(counts), [2, 3, 4])
        self.assertEqual(counts[4]["duplicate_of_input_row"], 2)
        self.assertEqual(counts[4]["same_product_count"], counts[2]["same_product_count"])
        self.assertIn("第 2 行", counts[4]["match_reason"])

    def test_bad_main_image_url_falls_back_to_product_page(self) -> None:
        products = self.write_products(["B012345678,https://example.invalid/dead.jpg"], header="ASIN,主图URL")
        runtime = self.runtime(products)

        def emb(_runtime, ref):
            return [1.0, 0.0] if ref.startswith("data:") else [0.0, 1.0]

        patches = [p for p in self.loop_patches() if getattr(p, "attribute", "") != "resolve_source_image"] + [
            patch.object(image, "call_multimodal_embedding", side_effect=emb),
            patch.object(
                image,
                "download_image",
                side_effect=[image.SourceImageDownloadError("HTTP 404"), self.source_image],
            ),
            patch.object(
                image,
                "load_source_product_main_image",
                return_value="https://m.media-amazon.com/images/I/real.jpg",
            ),
        ]
        started = [p.start() for p in patches]
        try:
            with redirect_stdout(io.StringIO()):
                self.assertEqual(image.run_image_competitor_crawl(runtime, dry_run=False), 0)
        finally:
            for p in reversed(patches):
                p.stop()
        started[-1].assert_called_once()
        job_dir = self.job_dir(runtime)
        reasons = [row["reason"] for row in image.read_jsonl(job_dir / "failures.jsonl")]
        self.assertIn("source_image_download_failed", reasons)
        self.assertEqual(len(image.read_jsonl(job_dir / "counts.jsonl")), 1)

    def test_download_rejects_non_image_content_type(self) -> None:
        response = MagicMock()
        response.raise_for_status.return_value = None
        response.headers = {"content-type": "text/html; charset=utf-8"}
        response.content = b"<html>"
        with patch.object(image.requests, "get", return_value=response):
            with self.assertRaises(image.SourceImageDownloadError):
                image.download_image("https://example.invalid/a.jpg", self.root / "img", "s")

    def test_main_exit_codes_run_summary_partial_workbook_and_job_label(self) -> None:
        products = self.write_products(["B0AAAAAAAA,,,,x", "B0BBBBBBBB,,,,x"])
        config_path = self.root / "config.json"
        config_path.write_text(json.dumps(self.raw_config(products)), encoding="utf-8")

        def emb(_runtime, ref):
            if "B0BBBBBBBB" in ref:
                raise image.EmbeddingProviderError("HTTP 503 after retries")
            return [1.0, 0.0] if ref.startswith("data:") else [0.0, 1.0]

        def merge(_d, _r, current, _s):
            row = {**candidate(1, source_asin=current["source_asin"]), "source_id": current["source_id"]}
            row["candidate_image_url"] = f"https://m.media-amazon.com/images/I/{current['source_asin']}.jpg"
            return [row]

        acquired_labels = []
        real_acquire = image.LocalSafetyController.acquire

        def recording_acquire(controller):
            acquired_labels.append(controller.job_label)
            return real_acquire(controller)

        patches = self.loop_patches(merge) + [
            patch.object(image, "call_multimodal_embedding", side_effect=emb),
            patch.object(image.LocalSafetyController, "acquire", recording_acquire),
        ]
        for p in patches:
            p.start()
        try:
            out = io.StringIO()
            with redirect_stdout(out), patch("sys.stderr", io.StringIO()):
                code = image.main(["--config", str(config_path)])
        finally:
            for p in reversed(patches):
                p.stop()
        self.assertEqual(code, run_outcome.EXIT_RETRY_LATER)
        job_dir = self.root / "outputs" / "round1-image"
        summary = json.loads((job_dir / "run_summary.json").read_text(encoding="utf-8"))
        self.assertEqual(summary["status"], "retry_later")
        self.assertEqual(summary["exit_code"], 20)
        workbooks = list(job_dir.glob("*.xlsx"))
        self.assertEqual(len(workbooks), 1)
        from openpyxl import load_workbook

        ws = load_workbook(workbooks[0]).active
        headers = [cell.value for cell in ws[1]]
        count_col = headers.index(image.COUNT_COLUMN_HEADER) + 1
        self.assertEqual(ws.cell(row=2, column=count_col).value, 0)
        self.assertIsNone(ws.cell(row=3, column=count_col).value)
        self.assertEqual(acquired_labels, ["round1-image"])
        # The holder record is cleared on release so a stale pid is never read as a running crawl.
        lock_file = Path(os.environ["HOME"]) / ".lc-amazon-data-crawl" / "safety" / "crawler.lock"
        self.assertEqual(lock_file.read_text(encoding="utf-8"), "")

    def test_main_config_errors_and_dry_run(self) -> None:
        with patch("sys.stderr", io.StringIO()), redirect_stdout(io.StringIO()):
            self.assertEqual(image.main(["--config", str(self.root / "missing.json")]), run_outcome.EXIT_CONFIG_ERROR)
            bad = self.root / "bad.json"
            bad.write_text("{not json", encoding="utf-8")
            self.assertEqual(image.main(["--config", str(bad)]), run_outcome.EXIT_CONFIG_ERROR)
            products = self.write_products(["B012345678,,,,x"])
            config_path = self.root / "config.json"
            config_path.write_text(json.dumps(self.raw_config(products)), encoding="utf-8")
            self.assertEqual(image.main(["--config", str(config_path), "--dry-run"]), 0)
            self.mini_config.write_text(json.dumps({"api_key": ""}), encoding="utf-8")
            self.assertEqual(
                image.main(["--config", str(config_path), "--dry-run"]),
                run_outcome.EXIT_CONFIG_ERROR,
            )
        self.assertFalse((self.root / "outputs" / "round1-image" / "run_summary.json").exists())


if __name__ == "__main__":
    unittest.main()
