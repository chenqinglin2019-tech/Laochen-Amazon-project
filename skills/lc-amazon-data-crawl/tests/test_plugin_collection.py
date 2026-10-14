import datetime as dt
import contextlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from amazon_plugin_crawler import (PanelState, PluginCrawler, fingerprint, limit_value,
                                   source_task, validate_config, build_tasks, connect_work_page)
from opencli_bridge import OpenCliBridge, decode_json
from run_outcome import CrawlStop
from safety_control import LocalSafetyController, RiskSignal, SafetyPausedError
from sellersprite_export import parse_export, verify_export, apply_filters, deduplicate, write_outputs
from openpyxl import Workbook, load_workbook


class Clock:
    def __init__(self):
        self.value = 1800000000.0
    def now(self):
        return self.value
    def sleep(self, seconds):
        self.value += seconds


class FakeSafety:
    def __init__(self, root):
        self.storage_root = root
        self.actions = []
    def heartbeat(self, *args, **kwargs): pass
    def note_risk_checked(self): pass
    def complete_review_success(self): pass
    def set_work_key(self, value): pass
    def before_remote_action(self, action): self.actions.append(action)
    def trip(self, *args, **kwargs): raise RuntimeError("risk")


class FakeBridge:
    def __init__(self, count=16, increments=None):
        self.count = count
        self.increments = list(increments or [])
        self.document_id = "doc"
        self.actions = []
        self.url = "https://www.amazon.com/s?me=SELLER1"
    def panel(self, action="snapshot", **kwargs):
        if action == "snapshot":
            return {"count": self.count, "busy": False, "panel": True, "more": bool(self.increments),
                    "more_count": int(bool(self.increments)), "rows": [], "url": self.url}
        self.actions.append(action)
        if action == "load":
            self.count += self.increments.pop(0)
        return {"ready": True}


def native_file(path, count, *, duplicate=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "US"
    ws.append(["#", "ASIN", "商品标题", "月销量", "子体销量", "评分", "评分数", "价格", "配送方式", "SP广告", "小类BSR", "小类目"])
    for n in range(count):
        asin = f"B{(n % 3 if duplicate else n):09d}"
        ws.append([n+1, asin, f"Product {n}", 120, "50+", 4.8, 1300, 19.99, "FBA", "", 123, "Tools"])
    wb.create_sheet("Brands").append(["Brand", "Sales"])
    wb.save(path)
    wb.close()


class PluginCollectionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.clock = Clock()
    def tearDown(self): self.tmp.cleanup()
    def cfg(self, **kwargs):
        return validate_config({"mode": "storefront", "job_id": "job", **kwargs}, self.root)
    def state(self, cfg=None):
        cfg = cfg or self.cfg()
        cfg["job_dir"].mkdir(parents=True, exist_ok=True)
        return PanelState(cfg["job_dir"] / "state.json", fingerprint(cfg, []))
    def crawler(self, cfg, bridge):
        return PluginCrawler(cfg, self.state(cfg), bridge, FakeSafety(self.root), now=self.clock.now, sleep=self.clock.sleep)

    def test_defaults_overrides_and_invalid_caps(self):
        cfg = self.cfg()
        self.assertEqual((cfg["store_max_records"], cfg["keyword_max_records"]), (1008, 2000))
        self.assertIsNone(self.cfg(store_max_records="all")["store_max_records"])
        self.assertEqual(self.cfg(keyword_max_records=3000)["keyword_max_records"], 3000)
        for value in [0, -1, True, None, "1008", 1.5, "ALL"]:
            with self.subTest(value=value), self.assertRaises(CrawlStop):
                self.cfg(store_max_records=value)

    def test_old_config_rejected_before_browser(self):
        for old in [{"browser_backend": "cdp"}, {"max_pages_per_keyword": 7}, {"store_page_limit": 20}]:
            with self.assertRaises(CrawlStop): self.cfg(**old)

    def test_batch_same_keyword_multiple_sorts_shared_source(self):
        (self.root / "input.csv").write_text("keyword\ncelery seed\ncelery seed\n", encoding="utf-8")
        cfg = self.cfg(mode="keyword_search", entry_mode="batch", keywords_file="input.csv",
                       keyword_sort_orders=["Featured", "Newest Arrivals"])
        tasks = build_tasks(cfg)
        self.assertEqual(len(tasks), 2)
        self.assertEqual(tasks[0]["source"], tasks[1]["source"])
        self.assertNotEqual(tasks[0]["key"], tasks[1]["key"])

    def test_budget_survives_resume_and_replaces_cumulative_snapshot(self):
        cfg = self.cfg(); state = self.state(cfg)
        one = source_task("https://www.amazon.com/s?me=SELLER1", "Featured")
        two = source_task("https://www.amazon.com/s?me=SELLER1&s=date-desc-rank", "Newest Arrivals")
        state.commit(one, [{}]*800, {}, 1008, "checkpoint")
        state.commit(one, [{}]*960, {}, 1008, "source_exhausted")
        state = PanelState(state.path, fingerprint(cfg, []))
        self.assertEqual(state.budget(two, 1008), 48)
        state.commit(two, [{}]*96, {}, 1008, "limit_reached")
        self.assertEqual(sum(t["retained_count"] for t in state.data["tasks"].values()), 1008)
        self.assertEqual(state.budget(one, 1008), 960)
        with self.assertRaises(CrawlStop): PanelState(state.path, "changed-plan")

    def test_load_to_limit_overshoots_only_one_batch(self):
        cfg = self.cfg(); bridge = FakeBridge(960, [96, 96]); crawler = self.crawler(cfg, bridge)
        task = source_task(bridge.url)
        exports = []
        crawler.export = lambda task, snap: (exports.append(snap["count"]) or [{}]*snap["count"], {"plugin_count": snap["count"]})
        crawler.materialize = lambda: None
        crawler.run_task(task, prepared=True)
        result = crawler.state.data["tasks"][task["key"]]
        self.assertEqual(bridge.actions, ["load"])
        self.assertEqual(exports, [1056])
        self.assertEqual(result["retained_count"], 1008)
        self.assertEqual(result["finish_reason"], "limit_reached")

    def test_at_limit_and_small_store_do_not_load(self):
        for count, increments, reason in [(1200, [96], "limit_reached"), (12, [], "source_exhausted")]:
            with self.subTest(count=count):
                cfg = self.cfg(job_id=f"job-{count}"); b = FakeBridge(count, increments); c = self.crawler(cfg, b)
                task = source_task(b.url)
                c.export = lambda t, s: ([{}]*s["count"], {})
                c.materialize = lambda: None
                c.run_task(task, prepared=True)
                self.assertEqual(b.actions, [])
                self.assertEqual(c.state.data["tasks"][task["key"]]["finish_reason"], reason)

    def test_missing_store_skips_without_plugin_controls_and_continues(self):
        cfg = self.cfg(); b = FakeBridge(); original = b.panel
        def panel(action="snapshot", **kwargs):
            if action == "snapshot":
                return {**original(), "count": None, "panel": False, "amazon_no_results": True}
            return original(action, **kwargs)
        b.panel = panel
        c = self.crawler(cfg, b); task = source_task(b.url)
        c.run_task(task)
        self.assertEqual(b.actions, [])
        self.assertGreaterEqual(self.clock.value - 1800000000, 10)
        saved = PanelState(c.state.path, fingerprint(cfg, []))
        self.assertTrue(saved.data["tasks"][task["key"]]["complete"])
        self.assertEqual(saved.data["tasks"][task["key"]]["finish_reason"], "store_no_results")
        self.assertEqual(len(saved.data["skipped_items"]), 1)
        self.assertIsNone(saved.data["current_task"])
        b.panel = original; b.url = "https://www.amazon.com/s?me=SELLER2"
        c.export = lambda t, snap: ([{}]*snap["count"], {})
        c.materialize = lambda: None
        next_task = source_task(b.url)
        c.run_task(next_task, prepared=True)
        self.assertTrue(c.state.data["tasks"][next_task["key"]]["complete"])

    def test_transient_empty_message_and_keyword_do_not_skip(self):
        for keyword in (False, True):
            cfg = self.cfg(job_id=f"transient-{keyword}", mode="keyword_search" if keyword else "storefront")
            b = FakeBridge(); original = b.panel
            if keyword: b.url = "https://www.amazon.com/s?k=test"
            def panel(action="snapshot", **kwargs):
                snap = original(action, **kwargs)
                if action == "snapshot":
                    snap["amazon_no_results"] = keyword or self.clock.value < 1800000006
                return snap
            b.panel = panel; c = self.crawler(cfg, b)
            c.export = lambda t, snap: ([{}]*snap["count"], {})
            c.materialize = lambda: None
            c.run_task(source_task(b.url), prepared=True)
            self.assertFalse(c.state.data.get("skipped_items"))

    def test_captcha_takes_priority_over_missing_store(self):
        cfg = self.cfg(); b = FakeBridge(); original = b.panel
        b.panel = lambda *a, **kw: {**original(), "amazon_no_results": True, "captcha": True}
        c = self.crawler(cfg, b)
        with self.assertRaisesRegex(RuntimeError, "risk"):
            c.run_task(source_task(b.url))
        self.assertFalse(c.state.data.get("skipped_items"))
        self.assertEqual(b.actions, [])

    def test_keyword_default_2000(self):
        cfg = self.cfg(mode="keyword_search"); b = FakeBridge(1990, [50, 50]); b.url = "https://www.amazon.com/s?k=celery"
        c = self.crawler(cfg, b); task = source_task(b.url)
        c.export = lambda t, s: ([{}]*s["count"], {})
        c.materialize = lambda: None
        c.run_task(task, prepared=True)
        self.assertEqual(c.state.data["tasks"][task["key"]]["retained_count"], 2000)
        self.assertEqual(b.actions, ["load"])

    def test_midpoint_exports_only_every_100_loads(self):
        cfg = self.cfg(store_max_records="all"); b = FakeBridge(1, [1]*101); c = self.crawler(cfg, b)
        counts = []; task = source_task(b.url)
        c.export = lambda t, s: (counts.append(s["count"]) or [{}]*s["count"], {})
        c.materialize = lambda: None
        c.run_task(task, prepared=True)
        self.assertEqual(counts, [101, 102])
        self.assertEqual(c.state.data["tasks"][task["key"]]["retained_count"], 102)

    def test_still_busy_never_double_clicks(self):
        cfg = self.cfg(); b = FakeBridge(); c = self.crawler(cfg, b)
        b.panel = lambda *a, **kw: {"count": 16, "busy": True, "panel": True, "more": True, "rows": []}
        with self.assertRaises(CrawlStop): c.stable(after_count=16)
        self.assertEqual(b.actions, [])

    def test_native_export_2100_and_exact_aliases(self):
        file = self.root / "Products(2100)-US.xlsx"; native_file(file, 2100)
        rows, report = parse_export(file, "amazon.com")
        verify_export(rows, report, {"count": 2100, "rows": [{"index": "2100", "title": "Product 2099"}]})
        self.assertEqual(rows[0]["sales_30_days_parent"], 120)
        self.assertEqual(rows[0]["sales_30_days_child"], "50+")
        self.assertEqual(rows[0]["rating_value"], 4.8)
        self.assertEqual(rows[0]["review_count"], 1300)
        selected, _ = apply_filters(rows, {"require_subcategory_rank": True}, True)
        self.assertEqual(len(selected), 2100)
        with self.assertRaises(CrawlStop): verify_export(rows[:2000], report, {"count": 2100})
        with self.assertRaises(CrawlStop): verify_export(rows, report, {"count": 2100, "rows": [{"index": "1", "title": "Wrong download"}]})

    def test_truncate_before_dedup_and_filters_no_backfill(self):
        file = self.root / "raw.xlsx"; native_file(file, 1056, duplicate=True)
        rows, _ = parse_export(file, "amazon.com")
        self.assertEqual(len(deduplicate(rows[:1008])), 3)
        other = {**rows[0], "marketplace": "amazon.ca"}
        self.assertEqual(len(deduplicate(rows[:1008]+[other])), 4)
        with self.assertRaises(CrawlStop): apply_filters([{**rows[0], "fulfillment_method": ""}], {"excluded_fulfillment_methods": ["FBM"]}, True)
        with self.assertRaises(CrawlStop): apply_filters([{**rows[0], "is_sponsored": None}], {}, False)

    def test_workbook_preserves_new_fields_and_no_fake_search_rank(self):
        file = self.root / "raw.xlsx"; native_file(file, 4)
        rows, _ = parse_export(file, "amazon.com")
        for row in rows: row.update(source_type="storefront", store_url="https://www.amazon.com/s?me=SELLER1")
        quality = write_outputs(self.root, rows, [], {}, True)
        self.assertEqual(quality["unique_asins"], 4)
        wb = load_workbook(self.root / "dedup_total.xlsx", read_only=True)
        header, first, *_ = list(wb.active.values)
        self.assertEqual(first[header.index("价格")], 19.99)
        self.assertEqual(list(header), ["#", "ASIN", "商品标题", "月销量", "子体销量", "评分", "评分数", "价格", "配送方式", "SP广告", "小类BSR", "小类目"])
        self.assertNotIn("最佳排名", header)
        self.assertEqual(wb.sheetnames, ["US", "Brands"])
        self.assertTrue(quality["native_workbook"]["format_preserved"])
        wb.close()

    def test_global_scheduler_survives_new_controller_and_rests(self):
        def controller(mode="supervised"):
            obj = LocalSafetyController(self.root / "safety", clock=lambda: dt.datetime.fromtimestamp(self.clock.now()), mode=mode, traffic_profile="plugin")
            obj._wait_until = lambda deadline, reason: self.clock.sleep(max(0, deadline.timestamp()-self.clock.now()))
            return obj
        with patch("safety_control.random.uniform", side_effect=lambda a, b: a):
            start = self.clock.now()
            s = controller()
            for _ in range(20): s.before_remote_action("plugin:load")
            self.assertEqual(self.clock.now()-start, 19*20)
            before = self.clock.now()
            controller().before_remote_action("plugin:load")
            self.assertEqual(self.clock.now()-before, 60)
            self.assertEqual(controller()._load_traffic()["actions_count"], 21)

    def test_export_cooldown_resume_and_commit_acknowledgement(self):
        cfg = self.cfg(); b = FakeBridge(4); c = self.crawler(cfg, b); task = source_task(b.url)
        path = self.root / "Products(4)-US.xlsx"; native_file(path, 4)
        os.utime(path, (self.clock.now()+90, self.clock.now()+90))
        result = {"downloaded": True, "state": "complete", "filename": str(path), "id": 41,
                  "url": "blob:https://www.amazon.com/test"}
        waits = []
        b.wait_download = lambda *args: waits.append(args) or result
        c.export_state_path.write_text(json.dumps({"next_allowed_at": self.clock.now()+60}))
        rows, report = c.export(task, b.panel())
        self.assertGreaterEqual(self.clock.now(), 1800000060)
        self.assertEqual(b.actions.count("export"), 1)
        self.assertTrue(c._read_export_state()["pending"])
        # Process restart after file validation but before committing must reuse raw file.
        b.wait_download = lambda *args: self.fail("must reuse staged raw file")
        resumed = PluginCrawler(cfg, c.state, b, c.safety, now=self.clock.now, sleep=self.clock.sleep)
        rows2, report2 = resumed.export(task, b.panel())
        self.assertEqual(len(rows2), 4)
        self.assertEqual(b.actions.count("export"), 1)
        c.state.commit(task, rows2, report2, 1008, "source_exhausted")
        c.acknowledge_export(task, report2)
        self.assertNotIn("pending", c._read_export_state())
        self.assertEqual(c._read_export_state()["download_ids"], [41])

    def test_export_ignores_old_foreign_downloads_and_retains_unknown_attempt(self):
        cfg = self.cfg(); b = FakeBridge(4); c = self.crawler(cfg, b); task = source_task(b.url)
        path = self.root / "Products(4)-US.xlsx"; native_file(path, 4)
        os.utime(path, (self.clock.now()-3600, self.clock.now()-3600))
        def wait(*args):
            self.clock.sleep(40)
            return {"downloaded": True, "state": "complete", "filename": str(path), "id": 99,
                    "url": "blob:https://www.amazon.com/old"}
        b.wait_download = wait
        with self.assertRaises(CrawlStop): c.export(task, b.panel())
        self.assertEqual(b.actions.count("export"), 1)
        self.assertTrue(c._read_export_state().get("pending"))
        self.assertFalse(c.state.data["tasks"])
        # Unknown prior outcome is reconciled without another click, even after restart.
        with self.assertRaises(CrawlStop): c.export(task, b.panel())
        self.assertEqual(b.actions.count("export"), 1)

    def test_export_known_failure_retries_once_with_60_second_gap(self):
        cfg = self.cfg(); b = FakeBridge(4); c = self.crawler(cfg, b); task = source_task(b.url)
        clicked = []; old_panel = b.panel
        def panel(action="snapshot", **kwargs):
            if action == "export": clicked.append(self.clock.now())
            return old_panel(action, **kwargs)
        b.panel = panel
        b.wait_download = lambda *args: {"downloaded": False, "state": "interrupted", "id": 8,
             "filename": str(self.root / "Products(4)-US.xlsx"), "url": "blob:https://www.amazon.com/failed"}
        with self.assertRaises(CrawlStop): c.export(task, b.panel())
        self.assertEqual(len(clicked), 2)
        self.assertGreaterEqual(clicked[1]-clicked[0], 60)
        self.assertNotIn("pending", c._read_export_state())

    def test_commit_resets_midpoint_counter_atomically(self):
        cfg = self.cfg(); state = self.state(cfg); task = source_task("https://www.amazon.com/s?me=S1")
        state.data["current_task"] = {**task, "loads_since_export": 100, "load_pending": True}
        state.commit(task, [], {}, 1008, "checkpoint")
        reloaded = PanelState(state.path, fingerprint(cfg, []))
        self.assertEqual(reloaded.data["current_task"]["loads_since_export"], 0)
        self.assertFalse(reloaded.data["current_task"]["load_pending"])

    def test_currency_suffixed_exact_headers(self):
        path = self.root / "currency.xlsx"; wb = Workbook(); ws = wb.active
        ws.append(["ASIN", "商品标题", "价格($)", "月销售额($)", "子体销售额($)", "FBA($)"])
        ws.append(["B000000001", "title", 9.99, 100, 20, 3]); wb.save(path); wb.close()
        rows, _ = parse_export(path, "amazon.com")
        self.assertEqual([rows[0][k] for k in ("price", "sales_amount", "child_sales_amount", "fba_fee")], [9.99, 100, 20, 3])

    def test_long_and_unattended_rest_survive_restart(self):
        for mode, count, seconds in (("supervised", 20, 60), ("supervised", 100, 120), ("unattended", 10, 120), ("unattended", 100, 120)):
            with self.subTest(mode=mode):
                root = self.root / mode
                s = LocalSafetyController(root, clock=lambda: dt.datetime.fromtimestamp(self.clock.now()),
                                          mode=mode, traffic_profile="plugin")
                s._write_traffic({"actions_count": count})
                s._wait_until = lambda deadline, reason: self.clock.sleep(max(0, deadline.timestamp()-self.clock.now()))
                before = self.clock.now()
                with patch("safety_control.random.uniform", side_effect=lambda a,b:a): s.before_remote_action("plugin:load")
                self.assertEqual(self.clock.now()-before, seconds)
                self.assertEqual(s._load_traffic()["actions_count"], count+1)

    def test_plugin_rest_bounds_and_spacing(self):
        for mode, count, bounds in (("supervised", 20, (60, 90)), ("supervised", 100, (120, 180)),
                                    ("unattended", 10, (120, 180)), ("unattended", 100, (120, 180))):
            for upper in (False, True):
                with self.subTest(mode=mode, count=count, upper=upper):
                    root = self.root / f"{mode}-{count}-{upper}"
                    s = LocalSafetyController(root, clock=lambda: dt.datetime.fromtimestamp(self.clock.now()),
                                              mode=mode, traffic_profile="plugin")
                    s._write_traffic({"actions_count": count})
                    s._wait_until = lambda deadline, reason: self.clock.sleep(max(0, deadline.timestamp()-self.clock.now()))
                    before = self.clock.now()
                    with patch("safety_control.random.uniform", side_effect=lambda a,b:b if upper else a):
                        s.before_remote_action("plugin:load")
                    self.assertEqual(self.clock.now()-before, bounds[int(upper)])
                    spacing = s._load_traffic()["next_allowed_at"]
                    self.assertEqual(dt.datetime.fromisoformat(spacing).timestamp()-self.clock.now(),
                                     (30 if upper else 20) if mode == "supervised" else (75 if upper else 45))

    def test_plugin_captcha_requires_confirmation_but_allows_early_review(self):
        for mode in ("supervised", "unattended"):
            for platform in ("amazon", "sellersprite"):
                with self.subTest(mode=mode, platform=platform):
                    root = self.root / f"{mode}-{platform}"
                    def controller(profile="plugin"):
                        return LocalSafetyController(root, clock=lambda: dt.datetime.fromtimestamp(self.clock.now()),
                                                     mode=mode, traffic_profile=profile)
                    s = controller()
                    with self.assertRaises(SafetyPausedError):
                        s.trip(RiskSignal(platform, "captcha_or_robot_check", "captcha"))
                    pause = s._load_pause()
                    self.assertEqual((dt.datetime.fromisoformat(pause["not_before"])-dt.datetime.fromtimestamp(self.clock.now())).total_seconds(), 86400)
                    with self.assertRaises(SafetyPausedError): controller().begin()
                    with self.assertRaises(SafetyPausedError): controller("default").begin(resume_after_review=True)
                    resumed = controller()
                    resumed.begin(resume_after_review=True)
                    resumed.before_remote_action("plugin:load")
                    self.assertTrue(resumed.pause_path.exists())
                    resumed.note_risk_checked()
                    resumed.complete_review_success()
                    self.assertFalse(resumed.pause_path.exists())
                    with self.assertRaises(SafetyPausedError):
                        resumed.trip(RiskSignal(platform, "captcha_or_robot_check", "captcha again"))
                    self.assertTrue(resumed.pause_path.exists())

    def test_plugin_other_risks_cannot_resume_before_deadline(self):
        for reason in ("http_429", "rate_limited", "account_restricted", "http_403", "verification_required"):
            with self.subTest(reason=reason):
                root = self.root / reason
                s = LocalSafetyController(root, clock=lambda: dt.datetime.fromtimestamp(self.clock.now()), traffic_profile="plugin")
                with self.assertRaises(SafetyPausedError): s.trip(RiskSignal("sellersprite", reason, reason))
                with self.assertRaises(SafetyPausedError): s.begin(resume_after_review=True)

    def test_plugin_interrupted_rest_preserves_deadline(self):
        root = self.root / "interrupted"
        def controller():
            return LocalSafetyController(root, clock=lambda: dt.datetime.fromtimestamp(self.clock.now()), traffic_profile="plugin")
        s = controller(); s._write_traffic({"actions_count": 20})
        def interrupt(deadline, reason):
            self.clock.sleep(30)
            raise KeyboardInterrupt
        s._wait_until = interrupt
        with patch("safety_control.random.uniform", side_effect=lambda a,b:b):
            with self.assertRaises(KeyboardInterrupt): s.before_remote_action("plugin:load")
        deadline = s._load_traffic()["rest_until"]
        resumed = controller()
        resumed._wait_until = lambda until, reason: self.clock.sleep(max(0, until.timestamp()-self.clock.now()))
        before = self.clock.now()
        with patch("safety_control.random.uniform", side_effect=lambda a,b:a): resumed.before_remote_action("plugin:load")
        self.assertEqual(self.clock.now()-before, 60)
        self.assertEqual(self.clock.now(), dt.datetime.fromisoformat(deadline).timestamp())

    def test_new_front_routes_never_start_testing_browser(self):
        from start_cdp_browser import should_auto_start
        from run_amazon_front_crawl import build_command
        for mode in ("keyword_search", "storefront"):
            for backend in ("opencli", "cdp"):
                raw = {"mode": mode, "browser_backend": backend, "browser_mode": "reuse"}
                self.assertFalse(should_auto_start(raw))
                path = self.root / "config.json"; path.write_text(json.dumps(raw))
                self.assertTrue(any(v.endswith("amazon_plugin_crawler.py") for v in build_command(path, True, False)))
                from amazon_front_crawler import cli_main
                with patch("amazon_plugin_crawler.main", return_value=40) as plugin:
                    self.assertEqual(cli_main(["--config", str(path), "--dry-run"]), 40)
                    plugin.assert_called_once()

    @unittest.skipUnless(shutil.which("node"), "Node is required for OpenCLI DOM fixtures")
    def test_public_dom_contract_without_browser(self):
        subprocess.run(["node", str(ROOT / "tests/test_sellersprite_dom.cjs")], check=True, capture_output=True)

    def test_keyword_native_filename_and_incomplete_export_report(self):
        cfg = self.cfg(mode="keyword_search"); b = FakeBridge(2100); b.url = "https://www.amazon.com/s?k=diamond+painting"
        c = self.crawler(cfg, b); task = source_task(b.url)
        path = self.root / "Search(diamond-painting)-2100-US.xlsx"; native_file(path, 2000)
        os.utime(path, (self.clock.now()+20, self.clock.now()+20))
        def downloaded(pattern, *args):
            self.assertEqual(pattern, "Search(diamond-painting)-2100-US")
            return {"id": 72, "downloaded": True, "state": "complete", "filename": str(path), "url": "blob:https://www.amazon.com/export"}
        b.wait_download = downloaded
        with self.assertRaises(CrawlStop): c.export(task, b.panel())
        c.materialize()
        report = json.loads((cfg["job_dir"] / "quality_report.json").read_text())
        self.assertEqual(report["unverified_exports"][0]["exported_count"], 2000)
        self.assertEqual(report["unverified_exports"][0]["plugin_count"], 2100)
        self.assertFalse(c.state.data["tasks"])
        self.assertTrue(c._read_export_state().get("pending"))

    def test_page_loss_replaces_snapshot_but_keeps_source_load_counter(self):
        cfg = self.cfg(); b = FakeBridge(16); c = self.crawler(cfg, b); task = source_task(b.url)
        b.session = "session"; b.target = "fixed"; b.expected_url = ""
        b.validate_target = lambda: {"url": b.url, "page": "fixed"}
        c.state.data["browser"] = {"url": b.url, "target": "fixed"}
        c.state.data["current_task"] = {**task, "document_id": "old", "plugin_count": 800}
        c.state.data["sources"][task["source"]] = {"loads_since_export": 37}
        connect_work_page(c)
        self.assertIsNone(c.state.data["current_task"])
        self.assertEqual(c.state.data["sources"][task["source"]]["loads_since_export"], 37)
        self.assertEqual(b.target, "fixed")
        self.assertEqual(b.actions, ["mark"])

    def test_changed_work_url_stops_without_rebind_or_navigation(self):
        cfg = self.cfg(); b = FakeBridge(); c = self.crawler(cfg, b)
        b.target = "fixed"; b.validate_target = lambda: {"url": "https://www.amazon.com/s?k=other", "page": "fixed"}
        c.state.data["browser"] = {"url": b.url, "target": "fixed"}
        with self.assertRaises(CrawlStop): connect_work_page(c)
        self.assertEqual(b.actions, [])

    def test_100th_load_finishing_source_exports_once(self):
        cfg = self.cfg(store_max_records="all"); b = FakeBridge(1, [1]*100); c = self.crawler(cfg, b)
        exports = []; task = source_task(b.url)
        c.export = lambda t,s: (exports.append(s["count"]) or [{}]*s["count"], {})
        c.materialize = lambda: None
        c.run_task(task, prepared=True)
        self.assertEqual(exports, [101])

    def test_readiness_check_preserves_existing_run_report(self):
        from amazon_plugin_crawler import run
        cfg = self.cfg(); self.state(cfg)
        report = cfg["job_dir"] / "run_summary.json"
        report.write_text('{"status":"retry_later","retained_count":42}')
        before = report.read_text(); b = FakeBridge(4)
        with patch("amazon_plugin_crawler.OpenCliBridge", return_value=b), \
                patch("amazon_plugin_crawler.connect_work_page", return_value={"url": b.url}), \
                patch("amazon_plugin_crawler.LocalSafetyController", return_value=LocalSafetyController(self.root / "check-safety")), \
                contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(run(cfg, [], check_only=True), 0)
        self.assertEqual(report.read_text(), before)
        self.assertEqual(b.actions, [])

    def test_bridge_never_uses_open_or_select(self):
        b = OpenCliBridge("test", "fixed", command=["opencli"])
        calls = []
        b.panel = lambda action, **kw: calls.append((action, kw)) or {}
        b.navigate("https://www.amazon.com/s?k=celery")
        self.assertEqual(calls[0][0], "navigate")
        self.assertEqual(b.target, "fixed")
        self.assertEqual(decode_json('Update notice\n{"count": 1008}\nnotice')["count"], 1008)


if __name__ == "__main__": unittest.main()
