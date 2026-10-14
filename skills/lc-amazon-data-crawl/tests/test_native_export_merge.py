import copy
import datetime as dt
import struct
import zlib
import tempfile
import unittest
import zipfile
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))

from lxml import etree as E
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from native_export_merge import merge_native_exports, NS, M, R, P, D, parse, encode, relpath, resolve
from sellersprite_export import parse_export, write_outputs
from run_outcome import CrawlStop
from merge_plugin_outputs import collect_committed


def fixture(path, start=0, changed_header=False):
    wb = Workbook(); ws = wb.active; ws.title = 'US'
    ws.append(['#', '图片', 'ASIN', '商品标题', '价格', '配送方式', '未知原始字段', '日期'])
    if changed_header:
        ws['G1'] = '另一字段'
    for c in ws[1]:
        c.fill = PatternFill('solid', fgColor='EE8A00')
        c.font = Font(name='Arial', bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30
    ws.column_dimensions['D'].width = 57
    ws.column_dimensions['G'].width = 31
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:H4'
    for i in range(3):
        ws.append([i+1, '', f'B{start+i:09}', f'Product {start+i}', 12.34+i,
                   'FBM' if i == 1 else 'FBA', f'原字段 {start+i}', dt.datetime(2026, 10, i+1)])
        ws.row_dimensions[i+2].height = 24+i
        ws.cell(i+2, 5).number_format = '$#,##0.00'
        ws.cell(i+2, 5).font = Font(name='Arial', color='FF0000' if start else '0000FF')
        ws.cell(i+2, 7).fill = PatternFill('solid', fgColor='AAFFAA')
        ws.cell(i+2, 3).hyperlink = f'https://amazon.com/dp/B{start+i:09}'
    wb.create_sheet('Brands').append(['品牌', '销量'])
    wb['Brands'].append(['Generic', start+1])
    wb.create_sheet('Sellers').append(['卖家', '销量'])
    wb['Sellers'].append([f'Seller{start}', 3])
    wb.create_sheet('Note')['A1'] = '原始说明'
    wb.save(path); wb.close()
    # Native OOXML fixtures exercise preservation without requiring Pillow.
    with zipfile.ZipFile(path) as z: parts={n:z.read(n) for n in z.namelist()}
    ws=parse(parts['xl/worksheets/sheet1.xml'])
    E.SubElement(ws,f'{{{M}}}drawing',{'{'+R+'}id':'rIdImages'})
    rels=parse(parts['xl/worksheets/_rels/sheet1.xml.rels'])
    E.SubElement(rels,f'{{{P}}}Relationship',Id='rIdImages',Type=R+'/drawing',Target='../drawings/drawing1.xml')
    drawing=E.Element(f'{{{D}}}wsDr',nsmap={'xdr':D,'a':'http://schemas.openxmlformats.org/drawingml/2006/main'})
    image_rels=E.Element(f'{{{P}}}Relationships',nsmap={None:P})
    for i in range(3):
        anchor=E.SubElement(drawing,f'{{{D}}}twoCellAnchor')
        for side in ('from','to'):
            marker=E.SubElement(anchor,f'{{{D}}}{side}')
            for name,value in [('col',1),('colOff',0),('row',i+1),('rowOff',0 if side=='from' else 100000)]:
                E.SubElement(marker,f'{{{D}}}{name}').text=str(value)
        pic=E.SubElement(anchor,f'{{{D}}}pic')
        nv=E.SubElement(pic,f'{{{D}}}nvPicPr');E.SubElement(nv,f'{{{D}}}cNvPr',id=str(i+1),name=f'Picture {i+1}')
        fill=E.SubElement(pic,f'{{{D}}}blipFill')
        E.SubElement(fill,'{http://schemas.openxmlformats.org/drawingml/2006/main}blip',{'{'+R+'}embed':f'rId{i+1}'})
        E.SubElement(anchor,f'{{{D}}}clientData')
        E.SubElement(image_rels,f'{{{P}}}Relationship',Id=f'rId{i+1}',Type=R+'/image',Target=f'../media/image{i}.png')
        def chunk(tag,data):return struct.pack('!I',len(data))+tag+data+struct.pack('!I',zlib.crc32(tag+data)&0xffffffff)
        parts[f'xl/media/image{i}.png']=b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR',struct.pack('!IIBBBBB',1,1,8,2,0,0,0))+chunk(b'IDAT',zlib.compress(bytes([0,start%255,i*50,180])))+chunk(b'IEND',b'')
    parts.update({'xl/worksheets/sheet1.xml':encode(ws),'xl/worksheets/_rels/sheet1.xml.rels':encode(rels),
                  'xl/drawings/drawing1.xml':encode(drawing),'xl/drawings/_rels/drawing1.xml.rels':encode(image_rels)})
    content=parse(parts['[Content_Types].xml']);ct='http://schemas.openxmlformats.org/package/2006/content-types'
    E.SubElement(content,f'{{{ct}}}Default',Extension='png',ContentType='image/png')
    E.SubElement(content,f'{{{ct}}}Override',PartName='/xl/drawings/drawing1.xml',ContentType='application/vnd.openxmlformats-officedocument.drawing+xml')
    parts['[Content_Types].xml']=encode(content)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,data in parts.items():z.writestr(n,data)


def images(path):
    with zipfile.ZipFile(path) as z:
        rels={r.get('Id'):r.get('Target') for r in parse(z.read('xl/drawings/_rels/drawing1.xml.rels'))}
        result=[]
        for anchor in parse(z.read('xl/drawings/drawing1.xml')):
            rid=anchor.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip').get('{'+R+'}embed')
            result.append((int(anchor.findtext('d:from/d:row',namespaces=NS)),z.read(resolve('xl/drawings/drawing1.xml',rels[rid]))))
        return result


class NativeExportMergeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.a = self.root/'a.xlsx'; self.b = self.root/'b.xlsx'
        fixture(self.a); fixture(self.b, 10)
    def tearDown(self): self.tmp.cleanup()

    def test_all_columns_styles_dimensions_images_and_links_survive_selected_rows(self):
        out = self.root/'out.xlsx'
        result = merge_native_exports(out, [{'path':self.a,'sheet':'US','rows':[2,4]},
                                            {'path':self.b,'sheet':'US','rows':[3]}])
        self.assertEqual((result['product_rows'],result['original_columns'],result['images']), (3,8,3))
        wb = load_workbook(out); a = load_workbook(self.a); b = load_workbook(self.b)
        try:
            self.assertEqual(wb.sheetnames, a.sheetnames)
            ws = wb['US']; self.assertEqual(ws.freeze_panes, 'A2')
            self.assertEqual(ws.auto_filter.ref, 'A1:H4')
            self.assertEqual(ws.column_dimensions['D'].width, 57)
            for row,(original,num) in enumerate([(a,2),(a,4),(b,3)],2):
                for col in range(1,9):
                    x,y = ws.cell(row,col), original['US'].cell(num,col)
                    self.assertEqual(x.value,y.value)
                    self.assertEqual(copy.copy(x.font),copy.copy(y.font))
                    self.assertEqual(copy.copy(x.fill),copy.copy(y.fill))
                    self.assertEqual(copy.copy(x.alignment),copy.copy(y.alignment))
                    self.assertEqual(x.number_format,y.number_format)
                self.assertEqual(ws.row_dimensions[row].height,original['US'].row_dimensions[num].height)
                self.assertEqual(ws.cell(row,3).hyperlink.target,original['US'].cell(num,3).hyperlink.target)
            actual=images(out)
            self.assertEqual([im[0] for im in actual],[1,2,3])
            self.assertEqual([im[1] for im in actual],[images(self.a)[0][1],images(self.a)[2][1],images(self.b)[1][1]])
            self.assertEqual(wb['Brands'].max_row,3)
        finally:
            wb.close();a.close();b.close()

    def test_pipeline_filters_and_dedup_select_original_first_row(self):
        rows,_ = parse_export(self.a,'amazon.com')
        # Repeated rows consume the input budget but must not duplicate pictures.
        quality = write_outputs(self.root, rows+[copy.deepcopy(rows[0])], [],
                                {'excluded_fulfillment_methods':['FBM']}, True)
        self.assertEqual(quality['filtered_count'],2)
        self.assertEqual(quality['native_workbook']['images'],2)
        wb = load_workbook(self.root/'dedup_total.xlsx')
        self.assertEqual(wb['US'].max_column,8)
        self.assertEqual(wb['US']['G3'].value,'原字段 2')
        self.assertEqual(len(images(self.root/'dedup_total.xlsx')),2); wb.close()

    def test_legacy_checkpoint_recovers_exact_committed_file_not_intermediate(self):
        rows,report = parse_export(self.a,'amazon.com')
        for row in rows:
            for key in ('raw_export_file','raw_export_sheet','raw_export_row'): row.pop(key)
            row.update(source_type='storefront',store_url='url',sort_order='Newest Arrivals')
        source = {**report,'raw_file':str(self.a),'source_type':'storefront', 'store_url':'url','sort_order':'Newest Arrivals'}
        quality = write_outputs(self.root,rows,[source],{},True)
        self.assertEqual(quality['native_workbook']['product_rows'],3)
        self.assertEqual(quality['native_workbook']['source_files'],1)

    def test_changed_schema_or_unsupported_feature_fails_without_replacing_output(self):
        out = self.root/'out.xlsx'; out.write_bytes(b'existing output')
        fixture(self.b,10,changed_header=True)
        with self.assertRaisesRegex(CrawlStop,'列名'):
            merge_native_exports(out,[{'path':self.a,'sheet':'US','rows':[2]}, {'path':self.b,'sheet':'US','rows':[2]}])
        self.assertEqual(out.read_bytes(),b'existing output')
        fixture(self.b,10)
        wb=load_workbook(self.b);wb['US'].merge_cells('G2:H2');wb.save(self.b);wb.close()
        with self.assertRaisesRegex(CrawlStop,'尚未支持'):
            merge_native_exports(out,[{'path':self.a,'sheet':'US','rows':[2]}, {'path':self.b,'sheet':'US','rows':[2]}])
        self.assertEqual(out.read_bytes(),b'existing output')

    def test_empty_selection_retains_original_header_format(self):
        out=self.root/'empty.xlsx'
        result=merge_native_exports(out,[{'path':self.a,'sheet':'US','rows':[]}])
        self.assertEqual(result['images'],0)
        wb=load_workbook(out);self.assertEqual(wb['US'].max_row,1)
        self.assertEqual(wb['US']['A1'].fill.fgColor.rgb,'00EE8A00');wb.close()

    def test_missing_raw_file_never_falls_back_to_simplified_table(self):
        rows,_=parse_export(self.a,'amazon.com');self.a.unlink()
        with self.assertRaises(CrawlStop): write_outputs(self.root,rows,[],{},True)
        self.assertFalse((self.root/'dedup_total.xlsx').exists())

    def test_different_custom_number_formats_are_remapped(self):
        wb=load_workbook(self.b);wb['US']['E2'].number_format='0.0000%';wb.save(self.b);wb.close()
        # openpyxl without Pillow cannot re-save images, so this check uses a
        # workbook without drawing nodes consistently on both sides.
        for path in (self.a,self.b):
            with zipfile.ZipFile(path) as z: parts={n:z.read(n) for n in z.namelist()}
            ws=parse(parts['xl/worksheets/sheet1.xml'])
            for node in ws.findall(f'{{{M}}}drawing'):ws.remove(node)
            parts['xl/worksheets/sheet1.xml']=encode(ws)
            if 'xl/worksheets/_rels/sheet1.xml.rels' in parts:
                rels=parse(parts['xl/worksheets/_rels/sheet1.xml.rels'])
                for node in list(rels):
                    if node.get('Type')==R+'/drawing':rels.remove(node)
                parts['xl/worksheets/_rels/sheet1.xml.rels']=encode(rels)
            with zipfile.ZipFile(path,'w') as z:
                for n,data in parts.items():z.writestr(n,data)
        out=self.root/'formats.xlsx'
        merge_native_exports(out,[{'path':self.a,'sheet':'US','rows':[2]}, {'path':self.b,'sheet':'US','rows':[2]}])
        wb=load_workbook(out);self.assertEqual(wb['US']['E2'].number_format,'$#,##0.00')
        self.assertEqual(wb['US']['E3'].number_format,'0.0000%');wb.close()

    def test_offline_merge_uses_only_committed_snapshots(self):
        import json
        job=self.root/'job';job.mkdir();(job/'raw_exports').mkdir()
        (job/'raw_exports/intermediate.xlsx').write_bytes(b'not committed')
        (job/'state.json').write_text(json.dumps({'tasks':{
            'one':{'raw_file':str(self.a),'records':[{'asin':'B000000000'}]},
            'skipped':{'raw_file':'','records':[]}}}))
        records,sources=collect_committed([job,job])
        self.assertEqual(len(records),1);self.assertEqual(len(sources),1)
        self.assertEqual(sources[0]['raw_file'],str(self.a))


if __name__ == '__main__': unittest.main()
