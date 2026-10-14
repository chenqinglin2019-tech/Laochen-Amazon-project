"""Parse native product exports, preserving raw files and exact field semantics."""
from __future__ import annotations

import datetime as dt
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from run_outcome import config_error, retry_later, replace_with_retry


ALIASES = {
    "asin": ("ASIN",), "parent_asin": ("父ASIN", "Parent ASIN"),
    "panel_index": ("#", "序号"), "title": ("商品标题", "产品标题", "Title", "Product Title"),
    "product_url": ("商品详情页链接", "商品URL", "Product URL", "Product Link"),
    "image_url": ("商品主图", "Image URL"), "brand_name": ("品牌", "Brand"),
    "category": ("大类目", "Category"), "bsr": ("大类BSR", "BSR"),
    "subcategory": ("小类目", "Subcategory", "Sub Category"),
    "subcategory_rank": ("小类BSR", "Subcategory BSR", "Sub Category BSR"),
    "sales_30_days_parent": ("月销量", "销量(父)", "近30天销量（父体）", "Monthly Sales", "Sales"),
    "sales_30_days_child": ("子体销量", "近30天销量（子体）", "Child Sales", "Variation Sales"),
    "sales_amount": ("月销售额", "销售额", "Monthly Revenue", "Revenue"),
    "child_sales_amount": ("子体销售额", "Child Revenue"),
    "price": ("价格", "Price"), "variation_count": ("变体数", "Variations"),
    "review_count": ("评分数", "评论数量", "Ratings", "Reviews"),
    "rating_value": ("评分", "评分值", "Rating"),
    "review_delta": ("月新增评分数", "月评新增", "Monthly New Ratings"),
    "fba_fee": ("FBA", "FBA费用", "FBA Fees"), "gross_margin": ("毛利率", "Gross Margin"),
    "fulfillment_method": ("配送方式", "配送", "Fulfillment", "Fulfillment Type"),
    "delivery_duration": ("配送时长", "Delivery Time"),
    "seller_name": ("Buybox卖家", "卖家名称", "Buybox Seller", "Buy Box Seller"),
    "seller_country": ("卖家所属地", "卖家所处国家", "Seller Country"),
    "launch_date": ("上架时间", "Date Available", "Available Date"),
    "organic_keywords_count": ("自然搜索词数量", "自然搜索词", "Organic Keywords"),
    "ad_keywords_count": ("广告搜索词数量", "广告搜索词", "Sponsored Keywords"),
    "sp_ad": ("SP广告", "SP", "Sponsored Products"),
    "brand_ad": ("品牌广告", "Sponsored Brands", "SB"),
}


def header_key(value: Any) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value or ""))).casefold()


# Native exports append marketplace currency to these exact labels.
for _field in ("sales_amount", "child_sales_amount", "price", "fba_fee"):
    ALIASES[_field] += tuple(f"{label}({currency})" for label in ALIASES[_field]
                             for currency in ("$", "€", "£", "¥", "￥", "₹", "USD", "CAD", "AUD", "MXN", "BRL", "EUR", "GBP", "JPY", "INR", "AED", "SAR", "PLN", "SEK", "SGD", "TRY"))

HEADER_FIELDS = {header_key(label): key for key, labels in ALIASES.items() for label in labels}
MISSING = {"", "/", "--", "-", "n/a", "暂无", "无数据"}


def cell_value(value: Any) -> Any:
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return "" if value.casefold() in MISSING else value
    return "" if value is None else value


def parse_export(path: Path, marketplace: str) -> tuple[list[dict], dict]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise retry_later("原始导出已保留，但 XLSX 无法解析；不会提交该文件。") from exc
    try:
        candidates = []
        for sheet in workbook.worksheets:
            if sheet.title.casefold() in {"brands", "sellers"}:
                continue
            for number, values in enumerate(sheet.iter_rows(min_row=1, max_row=15, values_only=True), 1):
                if any(header_key(v) == "asin" for v in values):
                    candidates.append((sheet, number, list(values)))
                    break
        if len(candidates) != 1:
            raise retry_later("导出中的产品主表无法唯一识别；请使用产品查询的普通导出。")
        sheet, header_row, headers = candidates[0]
        columns = {index: HEADER_FIELDS[header_key(label)] for index, label in enumerate(headers)
                   if header_key(label) in HEADER_FIELDS}
        fields = list(columns.values())
        if len(fields) != len(set(fields)):
            raise retry_later("导出存在语义重复的关键列，无法安全映射。")
        if "title" not in fields:
            raise retry_later("导出缺少商品标题列，请启用产品查询数据列。")
        rows, invalid = [], 0
        for excel_row, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            record = {name: cell_value(values[index]) for index, name in columns.items() if index < len(values)}
            asin = str(record.get("asin") or "").upper()
            if not re.fullmatch(r"[A-Z0-9]{10}", asin):
                # The vendor appends an explanatory footer, not a product.
                nonempty = [str(v) for v in values if v is not None]
                if nonempty and not any("导出时自动" in v or "duplicate ASIN" in v for v in nonempty):
                    invalid += 1
                continue
            record.update(asin=asin, marketplace=marketplace,
                          raw_export_file=str(path.resolve()), raw_export_sheet=sheet.title,
                          raw_export_row=excel_row,
                          raw_row={str(h): cell_value(v) for h, v in zip(headers, values) if h is not None})
            record["product_url"] = record.get("product_url") or f"https://{marketplace}/dp/{asin}"
            rank = str(record.pop("subcategory_rank", "") or "")
            category = str(record.pop("subcategory", "") or "")
            record["subcategory_bsr_ranks"] = [{"category_name": category, "rank": rank}] if rank else []
            record["subcategory_field_present"] = "subcategory_rank" in fields
            indicators = [record.get(name) for name in ("sp_ad", "brand_ad") if name in fields]
            record["is_sponsored"] = None if not indicators else any(
                str(v).strip().casefold() not in {"", "0", "false", "no", "否", "-", "/"} for v in indicators)
            rows.append(record)
        return rows, {"sheet": sheet.title, "columns": [str(v) for v in headers if v is not None],
                      "fields": fields, "invalid_rows": invalid, "exported_count": len(rows),
                      "unique_asins": len({row["asin"] for row in rows})}
    finally:
        workbook.close()


def verify_export(records: list[dict], report: dict, snapshot: dict) -> None:
    expected = snapshot.get("count")
    if expected is None or report["invalid_rows"] or len(records) != expected:
        raise retry_later(f"导出完整性未通过：插件 {expected} 条，文件 {len(records)} 条，"
                          f"异常行 {report['invalid_rows']}。原始文件已保留，可能存在截断、勾选、去重或展开变体。",
                          next_action="检查原始导出与插件设置，确认后恢复；当前结果为部分完成。")
    sample = {a for row in snapshot.get("rows", []) for a in row.get("asins", [])}
    exported = {row["asin"] for row in records}
    if not sample.issubset(exported):
        raise retry_later("导出缺少当前面板已显示的 ASIN，可能使用了旧排序快照；原始文件已保留。")
    # 5.0.6 renders product titles as spans without ASIN attributes/links.
    pairs = {(str(row.get("panel_index", "")), row.get("title", "")) for row in records}
    for row in snapshot.get("rows", []):
        if row.get("index") and row.get("title") and (row["index"], row["title"]) not in pairs:
            raise retry_later("导出与面板商品序号/标题不符，可能是其他下载或过期快照；未提交数据。")


def deduplicate(records: list[dict]) -> list[dict]:
    grouped: dict[tuple[str, str], dict] = {}
    for record in records:
        key = (record["marketplace"], record["asin"])
        if key not in grouped:
            grouped[key] = {**record, "sources": [], "appearance_count": 0}
        target = grouped[key]
        source = {k: record.get(k, "") for k in ("source_type", "keyword", "store_url", "store_name", "sort_order", "page_url")}
        if source not in target["sources"]:
            target["sources"].append(source)
        target["appearance_count"] += 1
        for field, value in record.items():
            if target.get(field) in (None, "") and value not in (None, ""):
                target[field] = value
    return list(grouped.values())


def apply_filters(records: list[dict], filters: dict, include_sponsored: bool) -> tuple[list[dict], dict]:
    # Reuse pure validation/parsing/filtering functions, not the old browser flow.
    from amazon_category_rank_crawler import build_product_filter_config, filter_product_records, normalize_fulfillment_method
    rules = build_product_filter_config({"product_filters": filters})
    fulfillment = bool(rules.allowed_fulfillment_methods or rules.excluded_fulfillment_methods or rules.allow_missing_fulfillment)
    for row in records:
        if fulfillment and not normalize_fulfillment_method(row.get("fulfillment_method")):
            if not (rules.allow_missing_fulfillment and row.get("fulfillment_method") == ""):
                raise config_error("已启用配送方式筛选，但导出有缺失或无法识别的配送方式；不将未知值视为通过。")
        if rules.require_subcategory_rank and not row.get("subcategory_field_present"):
            raise config_error("已启用子类目排名筛选，但导出缺少对应字段。")
        if not include_sponsored and row.get("is_sponsored") is None:
            raise config_error("已要求排除广告，但导出缺少广告标识字段。")
    selected, rejections = filter_product_records(records, rules)
    if not include_sponsored:
        rejections["sponsored"] = sum(bool(row.get("is_sponsored")) for row in selected)
        selected = [row for row in selected if not row.get("is_sponsored")]
    return selected, rejections


def write_outputs(job_dir: Path, records: list[dict], sources: list[dict], filters: dict,
                  include_sponsored: bool) -> dict:
    from native_export_merge import merge_native_exports
    merged = deduplicate(records)
    accepted, rejected = apply_filters(merged, filters, include_sponsored)
    temp = job_dir / ".records.jsonl.tmp"
    temp.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in records), encoding="utf-8")
    replace_with_retry(str(temp), job_dir / "records.jsonl")
    # The workbook uses original cells, not normalized JSON fields. JSON remains
    # the audit/filtering representation; never reconstruct the vendor workbook.
    exports = native_export_selection(records, accepted, sources)
    native_report = merge_native_exports(job_dir / "dedup_total.xlsx", exports) if exports else {
        "format_preserved": True, "product_rows": 0, "workbook": "", "message": "暂无已提交的原生导出"}
    missing = Counter(field for row in merged for field in ("title", "price", "sales_30_days_parent", "sales_30_days_child", "rating_value", "review_count") if row.get(field) in (None, ""))
    return {"retained_count": len(records), "unique_asins": len(merged), "filtered_count": len(accepted),
            "rejection_counts": rejected, "missing_fields": dict(missing), "sources": sources,
            "native_workbook": native_report}


def native_export_selection(records: list[dict], accepted: list[dict], sources: list[dict]) -> list[dict]:
    """Resolve exact original rows, including checkpoints made before provenance existed."""
    exports = {}
    for source in sources:
        if source.get("raw_file"):
            path = str(Path(source["raw_file"]).resolve())
            exports.setdefault(path, {"path": path, "sheet": source.get("sheet"), "rows": []})
    for row in records:
        if row.get("raw_export_file"):
            path = str(Path(row["raw_export_file"]).resolve())
            exports.setdefault(path, {"path": path, "sheet": row["raw_export_sheet"], "rows": []})
    legacy = {}
    for row in accepted:
        if row.get("raw_export_file"):
            path = str(Path(row["raw_export_file"]).resolve())
            original = row
        else:
            matches = []
            for source in sources:
                if not source.get("raw_file"):
                    continue
                if any(row.get(k, "") != source.get(k, "") for k in
                       ("source_type", "store_url", "keyword", "sort_order")):
                    continue
                candidate = str(Path(source["raw_file"]).resolve())
                if candidate not in legacy:
                    parsed, _ = parse_export(Path(candidate), row["marketplace"])
                    legacy[candidate] = parsed
                matches.extend((candidate, item) for item in legacy[candidate]
                               if item["asin"] == row["asin"] and item.get("raw_row") == row.get("raw_row"))
            if len(matches) != 1:
                raise retry_later("无法唯一定位原始下载表中的商品行；保留原文件，不生成简化合并表。")
            path, original = matches[0]
        if path not in exports:
            raise retry_later("合并商品未关联已提交的原始下载文件。")
        exports[path]["sheet"] = original["raw_export_sheet"]
        exports[path]["rows"].append(int(original["raw_export_row"]))
    return list(exports.values())
