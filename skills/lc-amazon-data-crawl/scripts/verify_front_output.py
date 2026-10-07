"""Verify one committed snapshot, source evidence and final JSONL/Excel."""
from __future__ import annotations

import argparse
import collections
import json
import re
from pathlib import Path

import amazon_front_crawler as front
import amazon_category_rank_crawler as shared
from browser_runtime import pid_is_running
from runtime_watchdog import write_json

SOURCE_PATTERNS = {
    "sales_30_days_parent": r"近30天销量\(父体\)\s*[:：]\s*([<>~≈]?\s*[^\s]+)",
    "sales_30_days_child": r"近30天销量\(子体\)\s*[:：]\s*([<>~≈]?\s*[^\s]+)",
    "fba_fee": r"FBA费用\s*[:：]\s*([^\s]+)",
    "gross_margin": r"毛利率\s*[:：]\s*([^\s]+)",
    "brand_name": r"品牌\s*[:：]\s*([^\r\n]*?)(?=\r?\n|\s+(?:卖家\s*[:：]|加入产品库|近30天销量|#[\d,]+\s+in\b)|$)",
    "seller_name": r"(?<![A-Za-z])卖家\s*[:：]\s*([^\r\n]*?)(?=\r?\n|\s+配送\s*[:：]|$)",
    "delivery_duration": r"(?<!Prime)配送时长\s*[:：]\s*([^\s]+)",
    "organic_keywords_count": r"自然搜索词\s*[:：]\s*([\d,]+)",
    "ad_keywords_count": r"广告(?:搜索|流量)词\s*[:：]\s*([\d,]+)",
    "fulfillment_method": r"配送\s*[:：]\s*(FBA|FBM|AMZ)",
    "launch_date": r"上架时间\s*[:：]\s*(\d{4}-\d{2}-\d{2}(?:\s*\([^)]+\))?|N/A|NA|--|—|–)",
}


def load(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def verify(job: Path, live=False):
    heartbeat = load(job / "run_heartbeat.json") if (job / "run_heartbeat.json").exists() else {}
    supervisor = load(job / "supervisor.json") if (job / "supervisor.json").exists() else {}
    if not live and supervisor.get("phase") not in {"finished", "interrupted", "needs_human", "risk_pause"} and pid_is_running(supervisor.get("pid", 0)):
        raise ValueError("监督运行尚未结束，当前请使用 --live。")
    if not live and heartbeat.get("phase") != "exited" and pid_is_running(heartbeat.get("pid", 0)):
        raise ValueError("采集仍在运行，最终验收请等退出，或使用 --live。")
    state = load(job / "state.json")
    order = list(state.get("completed_page_order") or state.get("completed_pages") or [])
    mismatches = []
    payloads = {}
    for path in sorted((job / "page_results").glob("*.json")):
        payload = load(path)
        key = payload.get("page_key")
        if key in payloads:
            mismatches.append({"field": "duplicate_page_commit", "page": key})
        payloads[key] = payload
    records = []
    for key in order:
        if key not in payloads:
            mismatches.append({"field": "missing_page_commit", "page": key})
        else:
            records.extend(payloads[key].get("records") or [])
    if not live and set(payloads) != set(order):
        mismatches.append({"field": "commit_manifest_mismatch"})
    identities = [(p, str(r.get("asin"))) for p in order if p in payloads for r in payloads[p].get("records") or []]
    if len(set(identities)) != len(identities):
        mismatches.append({"field": "duplicate_asin_in_batch"})
    if len(records) != int(state.get("records_count") or 0):
        mismatches.append({"field": "state_record_count", "expected": len(records), "actual": state.get("records_count")})
    checked = collections.Counter()
    absent = collections.Counter()
    unassessed = 0
    for row in records:
        text = row.get("_source_card_text") or ""
        if not re.fullmatch(r"[A-Z0-9]{10}", str(row.get("asin") or "")):
            mismatches.append({"field": "invalid_ASIN", "asin": row.get("asin")})
        row_checks = 0
        for field, pattern in SOURCE_PATTERNS.items():
            values = re.findall(pattern, text, re.I)
            if not values:
                absent[field] += 1
                continue
            expected = shared.normalize_space(values[-1])
            actual = shared.normalize_space(str(row.get(field) if row.get(field) is not None else ""))
            checked[field] += 1
            row_checks += 1
            if expected != actual:
                mismatches.append({"asin": row.get("asin"), "field": field, "expected": expected, "actual": actual})
        rating = re.findall(r"评分\(评分数\)\s*[:：]\s*(N/A|[\d.]+)\s*\(\s*(N/A|[\d,]+)\s*\)", text)
        if rating:
            for field, value in zip(("rating_value", "review_count"), rating[-1]):
                checked[field] += 1
                row_checks += 1
                if str(row.get(field) or "") != value:
                    mismatches.append({"asin": row.get("asin"), "field": field, "expected": value, "actual": row.get(field)})
        code = row.get("_source_country_flag_code")
        html_flags = set(re.findall(r"\bflag-icon-([a-z]{2})\b", row.get("_source_plugin_html") or "", re.I))
        country_evidence_missing = bool(html_flags and not code)
        if html_flags and code and str(code).lower() not in {flag.lower() for flag in html_flags}:
            mismatches.append({"asin": row.get("asin"), "field": "country_flag_vs_html", "actual": code})
        if code:
            checked["seller_country"] += 1
            row_checks += 1
            expected = shared.country_from_flag_code_or_text(code)
            if expected != (row.get("seller_country") or ""):
                mismatches.append({"asin": row.get("asin"), "field": "seller_country", "expected": expected, "actual": row.get("seller_country")})
        else:
            absent["seller_country"] += 1
        if not row_checks or not text or country_evidence_missing:
            unassessed += 1
    excel_checks = 0
    if not live:
        actual_records = front.read_jsonl(job / "records.jsonl")
        if actual_records != records:
            mismatches.append({"field": "commits_vs_jsonl", "expected": len(records), "actual": len(actual_records)})
        from openpyxl import load_workbook
        wb = load_workbook(job / "dedup_total.xlsx", read_only=True, data_only=True)
        try:
            values = list(wb.worksheets[0].iter_rows(values_only=True))
            headers = list(values[0]) if values else []
            if "ASIN" not in headers:
                mismatches.append({"field": "Excel_ASIN_header_missing"})
            else:
                column = headers.index("ASIN")
                rows = [r for r in values[1:] if any(v is not None for v in r)]
                asins = [str(r[column] or "") for r in rows]
                if len(asins) != len(set(asins)):
                    mismatches.append({"field": "Excel_duplicate_ASIN"})
                expected_rows = front.build_front_dedup_rows(records)
                if set(asins) != {r["asin"] for r in expected_rows} or len(rows) != len(expected_rows):
                    mismatches.append({"field": "Excel_ASIN_set_or_count"})
                by_asin = dict(zip(asins, rows))
                for row in expected_rows:
                    cells = by_asin.get(row["asin"])
                    if cells is None:
                        continue
                    for field in (*SOURCE_PATTERNS, "rating_value", "review_count", "seller_country", "subcategory_bsr_ranks"):
                        header = front.FRONT_FIELD_TO_HEADER[field]
                        if header not in headers:
                            mismatches.append({"field": "Excel_header_missing", "header": header})
                            continue
                        value = row.get(field)
                        expected = front.format_subcategory_bsr_ranks(value) if field == "subcategory_bsr_ranks" else ("" if value is None else str(value))
                        actual = cells[headers.index(header)]
                        actual = "" if actual is None else str(actual)
                        excel_checks += 1
                        if expected != actual:
                            mismatches.append({"asin": row["asin"], "field": "Excel:" + field, "expected": expected, "actual": actual})
        finally:
            wb.close()
    status = "failed" if mismatches else ("not_evaluated" if not records or unassessed else "passed")
    report = {"status": status, "mode": "live" if live else "final", "committed_batches": len(order),
              "raw_records": len(records), "unique_asins": len({r.get("asin") for r in records}),
              "source_field_checks": dict(checked), "source_fields_absent": dict(absent),
              "unassessed_records": unassessed, "excel_checks": excel_checks, "mismatches": mismatches[:200],
              "mismatch_count": len(mismatches), "note": "提取一致性验收；完整交付仍以质量与范围报告为准。"}
    write_json(job / ("verification_live.json" if live else "verification_final.json"), report)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--job", required=True, type=Path)
    parser.add_argument("--live", action="store_true")
    args = parser.parse_args()
    try:
        report = verify(args.job.resolve(), args.live)
    except (OSError, ValueError, KeyError) as exc:
        print(f"验收失败：{exc}")
        return 2
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
